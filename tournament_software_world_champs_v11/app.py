from flask import g, session, abort, Flask, render_template, request, redirect, url_for, send_file, flash, jsonify
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import Integer, String, Text, Boolean, ForeignKey
from sqlalchemy.orm import Mapped, mapped_column, relationship
import json, math, os, io, itertools, uuid
from collections import defaultdict
from datetime import datetime
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from werkzeug.utils import secure_filename

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get('DATA_DIR', BASE_DIR)
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, 'tournament.db')
APP_PORT = 5069
APP_URL = f'http://127.0.0.1:{APP_PORT}'
CREATOR_CREDIT = 'created by Michał Klimczak (klimczak.sc@gmail.com)'
UPLOAD_DIR = os.path.join(DATA_DIR, 'uploads')
os.makedirs(UPLOAD_DIR, exist_ok=True)

app = Flask(__name__)
app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{DB_PATH}'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY')
if not app.config['SECRET_KEY']:
    secret_path = os.path.join(DATA_DIR, '.session-secret')
    if not os.path.exists(secret_path):
        with open(secret_path, 'w') as f:
            f.write(uuid.uuid4().hex + uuid.uuid4().hex)
        os.chmod(secret_path, 0o600)
    with open(secret_path) as f:
        app.config['SECRET_KEY'] = f.read()
app.config.update(MAX_CONTENT_LENGTH=8 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE', '1' if os.environ.get('RAILWAY_ENVIRONMENT_ID') else '0') == '1',
    PERMANENT_SESSION_LIFETIME=__import__('datetime').timedelta(hours=12))
db = SQLAlchemy(app)

POPULAR_SYSTEMS = [
    ('group_knockout', 'Group Stage + Championship'),
    ('single_elimination', 'Single Elimination'),
    ('double_elimination', 'Double Elimination'),
    ('round_robin', 'Round Robin'),
    ('swiss', 'Swiss System'),
]

class Tournament(db.Model):
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    name: Mapped[str] = mapped_column(String(200), default='New Tournament')
    player_count: Mapped[int] = mapped_column(Integer, default=8)
    system: Mapped[str] = mapped_column(String(50), default='single_elimination')
    consolation: Mapped[bool] = mapped_column(Boolean, default=False)
    court_count: Mapped[int] = mapped_column(Integer, default=1)
    created_at: Mapped[str] = mapped_column(String(50), default=lambda: datetime.utcnow().isoformat())
    settings_json: Mapped[str] = mapped_column(Text, default='{}')
    players = relationship('Player', cascade='all, delete-orphan', backref='tournament', order_by='Player.seed')
    phases = relationship('Phase', cascade='all, delete-orphan', backref='tournament', order_by='Phase.order_index')
    matches = relationship('Match', cascade='all, delete-orphan', backref='tournament', order_by='Match.order_index')

class Player(db.Model):
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    tournament_id: Mapped[int] = mapped_column(ForeignKey('tournament.id'))
    name: Mapped[str] = mapped_column(String(200))
    club: Mapped[str] = mapped_column(String(200), default='')
    seed: Mapped[int] = mapped_column(Integer, default=1)

class Phase(db.Model):
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    tournament_id: Mapped[int] = mapped_column(ForeignKey('tournament.id'))
    name: Mapped[str] = mapped_column(String(200))
    phase_type: Mapped[str] = mapped_column(String(50), default='main')
    order_index: Mapped[int] = mapped_column(Integer, default=1)
    best_of: Mapped[int] = mapped_column(Integer, default=3)
    points_to: Mapped[int] = mapped_column(Integer, default=21)
    config_json: Mapped[str] = mapped_column(Text, default='{}')

class Match(db.Model):
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    tournament_id: Mapped[int] = mapped_column(ForeignKey('tournament.id'))
    phase_id: Mapped[int] = mapped_column(ForeignKey('phase.id'), nullable=True)
    round_name: Mapped[str] = mapped_column(String(100), default='Round')
    bracket: Mapped[str] = mapped_column(String(50), default='main')
    order_index: Mapped[int] = mapped_column(Integer, default=1)
    court_number: Mapped[int] = mapped_column(Integer, default=1)
    player1_id: Mapped[int] = mapped_column(ForeignKey('player.id'), nullable=True)
    player2_id: Mapped[int] = mapped_column(ForeignKey('player.id'), nullable=True)
    score_json: Mapped[str] = mapped_column(Text, default='[]')
    winner_id: Mapped[int] = mapped_column(ForeignKey('player.id'), nullable=True)
    loser_id: Mapped[int] = mapped_column(ForeignKey('player.id'), nullable=True)
    notes: Mapped[str] = mapped_column(Text, default='')

    player1 = relationship('Player', foreign_keys=[player1_id])
    player2 = relationship('Player', foreign_keys=[player2_id])
    winner = relationship('Player', foreign_keys=[winner_id])
    loser = relationship('Player', foreign_keys=[loser_id])
    phase = relationship('Phase')


def settings_for(t: Tournament):
    base = {'qualifiers_per_group': 2, 'best_next_place_qualifiers': 0, 'wildcard_comparison': 'ppm', 'group_count': 2, 'group_assignment_mode': 'auto', 'manual_groups': {}, 'pairing_mode': 'first_round_only', 'bye_policy': 'top_seeds', 'logo_filename': ''}
    try:
        base.update(json.loads(t.settings_json or '{}'))
    except Exception:
        pass
    return base


def save_settings_dict(t: Tournament, data: dict):
    t.settings_json = json.dumps(data, ensure_ascii=False)


def next_power_of_two(n):
    p = 1
    while p < n:
        p *= 2
    return p


def balanced_seed_positions(size):
    arr = [1, 2]
    while len(arr) < size:
        size2 = len(arr) * 2 + 1
        new = []
        for x in arr:
            new.extend([x, size2 - x])
        arr = new
    return arr


def seeded_bracket(players):
    n = len(players)
    size = next_power_of_two(max(2, n))
    pos = balanced_seed_positions(size)
    slots = [None] * size
    by_seed = sorted(players, key=lambda p: p.seed)
    for idx, p in enumerate(by_seed, start=1):
        if idx > size:
            break
        slots[pos.index(idx)] = p
    return slots


def label_for_places(start_place, end_place):
    if start_place == 1:
        return 'Championship'
    if start_place == 3 and end_place == 4:
        return '3rd Place'
    return f'Places {start_place}-{end_place}'


def round_label(section_label, round_index, total_rounds):
    if total_rounds == 1:
        return section_label
    names = {1: 'Final', 2: 'Semifinal'}
    if round_index == total_rounds:
        return f'{section_label} - Final'
    if round_index == total_rounds - 1 and total_rounds >= 2:
        return f'{section_label} - Semifinal'
    return f'{section_label} - Round {round_index + 1}'


def create_default_phases(t: Tournament):
    t.phases.clear()
    if t.system == 'group_knockout':
        phases = [
            Phase(name='Group Stage', phase_type='groups', order_index=1, best_of=3, points_to=21),
            Phase(name='Championship & Placement', phase_type='main', order_index=2, best_of=5, points_to=21),
        ]
    elif t.system == 'single_elimination':
        phases = [Phase(name='Championship & Placement', phase_type='main', order_index=1, best_of=5, points_to=21)]
    elif t.system == 'double_elimination':
        phases = [
            Phase(name='Winners Bracket', phase_type='main', order_index=1, best_of=5, points_to=21),
            Phase(name='Losers Bracket', phase_type='losers', order_index=2, best_of=3, points_to=21),
            Phase(name='Final', phase_type='final', order_index=3, best_of=5, points_to=21),
        ]
    elif t.system == 'round_robin':
        phases = [Phase(name='League', phase_type='groups', order_index=1, best_of=3, points_to=21)]
    else:
        phases = [
            Phase(name='Swiss', phase_type='swiss', order_index=1, best_of=3, points_to=21),
            Phase(name='Placement', phase_type='placement', order_index=2, best_of=5, points_to=21),
        ]
    for ph in phases:
        t.phases.append(ph)


def create_tournament(name, player_count, system, consolation, court_count):
    t = Tournament(name=name, player_count=player_count, system=system, consolation=consolation, court_count=court_count)
    save_settings_dict(t, {'qualifiers_per_group': 2, 'best_next_place_qualifiers': 0, 'group_count': 2, 'group_assignment_mode': 'auto', 'manual_groups': {}})
    for i in range(1, player_count + 1):
        t.players.append(Player(name=f'Player {i}', seed=i))
    create_default_phases(t)
    db.session.add(t)
    db.session.commit()
    return t


def clear_matches(tournament):
    Match.query.filter_by(tournament_id=tournament.id).delete()
    db.session.commit()


def meta_for(match: Match):
    try:
        return json.loads(match.notes or '{}')
    except Exception:
        return {}


def set_meta(match: Match, meta: dict):
    match.notes = json.dumps(meta, ensure_ascii=False)


def parse_score(score_text):
    sets = []
    winner = None
    if not score_text.strip():
        return sets, winner, None
    for raw in score_text.split(','):
        raw = raw.strip().replace(':', '-').replace('/', '-').replace(' ', '')
        if '-' not in raw:
            continue
        a, b = raw.split('-', 1)
        try:
            a, b = int(a), int(b)
        except Exception:
            continue
        sets.append([a, b])
    p1_sets = sum(1 for a, b in sets if a > b)
    p2_sets = sum(1 for a, b in sets if b > a)
    if p1_sets > p2_sets:
        winner = 1
    elif p2_sets > p1_sets:
        winner = 2
    return sets, winner, {'p1_sets': p1_sets, 'p2_sets': p2_sets}


def score_to_text(score_json):
    try:
        sets = json.loads(score_json or '[]')
    except Exception:
        sets = []
    return ', '.join(f'{a}-{b}' for a, b in sets)


def round_robin_pairings(players):
    arr = list(players)
    if len(arr) % 2 == 1:
        arr.append(None)
    n = len(arr)
    rounds = []
    for _ in range(n - 1):
        pairs = []
        for i in range(n // 2):
            a = arr[i]
            b = arr[n - 1 - i]
            if a and b:
                pairs.append((a, b))
        rounds.append(pairs)
        arr = [arr[0]] + [arr[-1]] + arr[1:-1]
    return rounds


def standings_for_matches(matches, players):
    table = {p.id: {'player': p, 'matches': 0, 'wins': 0, 'losses': 0, 'sets_won': 0, 'sets_lost': 0, 'points_won': 0, 'points_lost': 0} for p in players}
    for m in matches:
        sets = json.loads(m.score_json or '[]')
        if not (m.player1_id and m.player2_id and sets):
            continue
        p1 = table[m.player1_id]
        p2 = table[m.player2_id]
        s1 = sum(1 for a, b in sets if a > b)
        s2 = sum(1 for a, b in sets if b > a)
        p1['matches'] += 1
        p2['matches'] += 1
        p1['sets_won'] += s1
        p1['sets_lost'] += s2
        p2['sets_won'] += s2
        p2['sets_lost'] += s1
        p1['points_won'] += sum(a for a, b in sets)
        p1['points_lost'] += sum(b for a, b in sets)
        p2['points_won'] += sum(b for a, b in sets)
        p2['points_lost'] += sum(a for a, b in sets)
        if s1 > s2:
            p1['wins'] += 1
            p2['losses'] += 1
        elif s2 > s1:
            p2['wins'] += 1
            p1['losses'] += 1
    for row in table.values():
        row['points'] = row['wins'] * 2
        row['set_diff'] = row['sets_won'] - row['sets_lost']
        row['point_diff'] = row['points_won'] - row['points_lost']
    return sorted(table.values(), key=lambda r: (-r['points'], -r['wins'], -r['set_diff'], -r['point_diff'], r['player'].seed))


def compute_group_tables(t):
    groups = defaultdict(list)
    players_by_group = defaultdict(dict)
    all_matches = Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index).all()
    for m in all_matches:
        if m.bracket.startswith('group_'):
            groups[m.bracket].append(m)
            if m.player1:
                players_by_group[m.bracket][m.player1.id] = m.player1
            if m.player2:
                players_by_group[m.bracket][m.player2.id] = m.player2
    result = []
    for key in sorted(groups.keys()):
        rows = standings_for_matches(groups[key], list(players_by_group[key].values()))
        result.append((key, rows))
    settings = settings_for(t)
    position = max(1, int(settings.get('qualifiers_per_group', 2))) + 1
    if settings.get('wildcard_comparison') == 'exclude_bottom' and int(settings.get('best_next_place_qualifiers', 0)) > 0:
        eligible = [(key, rows) for key, rows in result if len(rows) >= position]
        target_size = min((len(rows) for _, rows in eligible), default=0)
        for key, rows in eligible:
            if len(rows) <= target_size:
                continue
            # Use FINAL full-table positions to choose the omitted opponents.
            # Do not reorder the original table or replace its displayed totals.
            retained = {row['player'].id for row in rows[:target_size]}
            counted = [m for m in groups[key] if m.player1_id in retained and m.player2_id in retained]
            adjusted = standings_for_matches(counted, [row['player'] for row in rows[:target_size]])
            candidate = rows[position - 1]
            candidate['wildcard_comparison'] = next(row for row in adjusted if row['player'].id == candidate['player'].id)
            candidate['excluded_opponents'] = [row['player'].name for row in rows[target_size:]]
    return result


def comparison_metrics(row):
    return row.get('wildcard_comparison', row)


def build_group_sources(group_tables, positions):
    bands = []
    tables_only = [rows for _, rows in group_tables]
    if len(tables_only) == 2 and len(positions) == 2:
        a, b = tables_only
        bands = [
            {'kind': 'group_pos', 'group': 'group_1', 'pos': positions[0]},
            {'kind': 'group_pos', 'group': 'group_2', 'pos': positions[1]},
            {'kind': 'group_pos', 'group': 'group_2', 'pos': positions[0]},
            {'kind': 'group_pos', 'group': 'group_1', 'pos': positions[1]},
        ]
    elif len(tables_only) >= 4 and len(positions) == 2:
        gnames = ['group_1', 'group_2', 'group_3', 'group_4']
        bands = [
            {'kind': 'group_pos', 'group': gnames[0], 'pos': positions[0]},
            {'kind': 'group_pos', 'group': gnames[3], 'pos': positions[1]},
            {'kind': 'group_pos', 'group': gnames[1], 'pos': positions[0]},
            {'kind': 'group_pos', 'group': gnames[2], 'pos': positions[1]},
            {'kind': 'group_pos', 'group': gnames[2], 'pos': positions[0]},
            {'kind': 'group_pos', 'group': gnames[1], 'pos': positions[1]},
            {'kind': 'group_pos', 'group': gnames[3], 'pos': positions[0]},
            {'kind': 'group_pos', 'group': gnames[0], 'pos': positions[1]},
        ]
    else:
        for gi, _ in enumerate(group_tables, start=1):
            for pos in positions:
                bands.append({'kind': 'group_pos', 'group': f'group_{gi}', 'pos': pos})
    return bands


def add_match(tournament_id, phase_id, order_index, bracket, round_name, source1, source2, section, place_start, place_end, depth, slot_index, key):
    m = Match(
        tournament_id=tournament_id,
        phase_id=phase_id,
        round_name=round_name,
        bracket=bracket,
        order_index=order_index,
        court_number=1,
    )
    set_meta(m, {
        'key': key,
        'sources': [source1, source2],
        'section': section,
        'place_start': place_start,
        'place_end': place_end,
        'depth': depth,
        'slot_index': slot_index,
    })
    db.session.add(m)
    return m


def build_exact_bracket(tournament_id, phase_id, sources, place_start, key_prefix, order_state, bracket='main', depth_base=0):
    sources = list(sources)
    n = len(sources)
    if n <= 1:
        return []

    section = label_for_places(place_start, place_start + n - 1)
    if n == 2:
        key = f'{key_prefix}_m1'
        order_state['value'] += 1
        add_match(
            tournament_id,
            phase_id,
            order_state['value'],
            bracket,
            section,
            sources[0],
            sources[1],
            section,
            place_start,
            place_start + 1,
            depth_base,
            0,
            key,
        )
        return [{'kind': 'winner', 'match_key': key}, {'kind': 'loser', 'match_key': key}]

    half = n // 2
    total_rounds = int(math.log2(n))
    winners = []
    losers = []
    round_name = round_label(section, depth_base, total_rounds + depth_base)
    for i in range(0, n, 2):
        key = f'{key_prefix}_r{depth_base+1}_m{i//2+1}'
        order_state['value'] += 1
        add_match(
            tournament_id,
            phase_id,
            order_state['value'],
            bracket,
            round_name,
            sources[i],
            sources[i+1],
            section,
            place_start,
            place_start + n - 1,
            depth_base,
            i // 2,
            key,
        )
        winners.append({'kind': 'winner', 'match_key': key})
        losers.append({'kind': 'loser', 'match_key': key})

    build_exact_bracket(tournament_id, phase_id, winners, place_start, f'{key_prefix}_W', order_state, bracket=bracket, depth_base=depth_base + 1)
    build_exact_bracket(
        tournament_id,
        phase_id,
        losers,
        place_start + half,
        f'{key_prefix}_L',
        order_state,
        bracket='placement' if place_start > 1 or bracket == 'placement' else 'main',
        depth_base=depth_base + 1,
    )
    return []


def build_exact_bracket_any(tournament_id, phase_id, sources, place_start, key_prefix, order_state, bracket='main', depth_base=0):
    """Build exact-placement brackets for any field size >= 2.

    Non-power-of-two fields use a preliminary/play-in round. Higher-priority
    sources receive byes; preliminary losers continue into a lower placement
    bracket so every source can still earn an exact place on court.
    """
    sources = [src for src in list(sources) if src]
    n = len(sources)
    if n <= 1:
        return []
    if n & (n - 1) == 0:
        return build_exact_bracket(tournament_id, phase_id, sources, place_start, key_prefix, order_state, bracket=bracket, depth_base=depth_base)

    core_size = 1 << (n.bit_length() - 1)
    play_in_matches = n - core_size
    bye_count = n - 2 * play_in_matches
    bye_sources = sources[:bye_count]
    prelim_sources = sources[bye_count:]
    prelim_winners = []
    prelim_losers = []
    section = label_for_places(place_start, place_start + n - 1)

    # Pair the strongest play-in source with the weakest, then work inward.
    # With 4 groups + 2/group + 2 best third-place teams this means the
    # wildcards face lower-priority automatic qualifiers rather than each other.
    for i in range(play_in_matches):
        src1 = prelim_sources[i]
        src2 = prelim_sources[-1 - i]
        key = f'{key_prefix}_playin_m{i+1}'
        order_state['value'] += 1
        add_match(
            tournament_id, phase_id, order_state['value'], bracket,
            f'{section} - Preliminary', src1, src2, section,
            place_start, place_start + n - 1, depth_base, i, key
        )
        prelim_winners.append({'kind': 'winner', 'match_key': key})
        prelim_losers.append({'kind': 'loser', 'match_key': key})

    core_sources = bye_sources + prelim_winners
    build_exact_bracket_any(
        tournament_id, phase_id, core_sources, place_start,
        f'{key_prefix}_CORE', order_state, bracket=bracket, depth_base=depth_base + 1
    )
    build_exact_bracket_any(
        tournament_id, phase_id, prelim_losers, place_start + core_size,
        f'{key_prefix}_PLAYOUT', order_state, bracket='placement', depth_base=depth_base + 1
    )
    return []


def build_championship_bracket_any(tournament_id, phase_id, sources, key_prefix, order_state, bracket='main', depth_base=0):
    """Championship-only bracket for any field size, using play-ins and byes."""
    sources = [src for src in list(sources) if src]
    n = len(sources)
    if n <= 1:
        return
    if n & (n - 1):
        core_size = 1 << (n.bit_length() - 1)
        play_in_matches = n - core_size
        bye_count = n - 2 * play_in_matches
        bye_sources = sources[:bye_count]
        prelim_sources = sources[bye_count:]
        winners = []
        for i in range(play_in_matches):
            key = f'{key_prefix}_playin_m{i+1}'
            order_state['value'] += 1
            add_match(
                tournament_id, phase_id, order_state['value'], bracket,
                'Championship - Preliminary', prelim_sources[i], prelim_sources[-1-i],
                'Championship', 1, n, depth_base, i, key
            )
            winners.append({'kind': 'winner', 'match_key': key})
        sources = bye_sources + winners
        depth_base += 1

    current = sources
    depth = depth_base
    while len(current) > 1:
        nxt = []
        round_count = int(math.log2(len(current)))
        rn = round_label('Championship', depth, round_count + depth)
        for i in range(0, len(current), 2):
            key = f'{key_prefix}_d{depth}_m{i//2+1}'
            order_state['value'] += 1
            add_match(
                tournament_id, phase_id, order_state['value'], bracket, rn,
                current[i], current[i+1], 'Championship', 1, n, depth, i//2, key
            )
            nxt.append({'kind': 'winner', 'match_key': key})
        current = nxt
        depth += 1


def best_group_position_rows(group_tables, position):
    """Rank teams occupying the same position across all groups.

    Tie-break order mirrors the group table: points, wins, set difference,
    point difference, then original seed.
    """
    candidates = []
    for group_key, rows in group_tables:
        if 1 <= position <= len(rows):
            row = rows[position - 1]
            candidates.append((group_key, row))
    candidates.sort(key=lambda item: (
        -(comparison_metrics(item[1])['points'] / max(1, comparison_metrics(item[1])['matches'])),
        -(comparison_metrics(item[1])['wins'] / max(1, comparison_metrics(item[1])['matches'])),
        -(comparison_metrics(item[1])['set_diff'] / max(1, comparison_metrics(item[1])['matches'])),
        -(comparison_metrics(item[1])['point_diff'] / max(1, comparison_metrics(item[1])['matches'])),
        item[1]['player'].seed
    ))
    return candidates



def championship_qualifier_rows(t):
    """Return the actual qualified teams after the full group stage.

    Each item contains player, source group and finishing position. Wildcards
    are ranked with the same cross-group tie-break used by resolve_source().
    """
    if t.system != 'group_knockout' or not group_stage_complete(t.id):
        return []
    settings = settings_for(t)
    q = max(1, int(settings.get('qualifiers_per_group', 2)))
    wild = max(0, int(settings.get('best_next_place_qualifiers', 0)))
    tables = compute_group_tables(t)
    out = []
    seen = set()
    for group_key, rows in tables:
        for pos in range(1, min(q, len(rows)) + 1):
            player = rows[pos-1]['player']
            if player.id not in seen:
                out.append({'player': player, 'group': group_key, 'position': pos, 'wildcard_rank': None})
                seen.add(player.id)
    if wild:
        next_pos = q + 1
        for rank, (group_key, row) in enumerate(best_group_position_rows(tables, next_pos)[:wild], start=1):
            player = row['player']
            if player.id not in seen:
                out.append({'player': player, 'group': group_key, 'position': next_pos, 'wildcard_rank': rank})
                seen.add(player.id)
    return out


def championship_entry_matches(t):
    """Matches whose slots are direct entrants from the group stage.

    For the World-Championship format currently used (8 qualifiers), these are
    the four quarterfinals. After a manual override the sources become explicit
    player sources, so the same matches remain editable.
    """
    candidates = []
    for m in Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index, Match.id).all():
        meta = meta_for(m)
        if meta.get('section') != 'Championship' or m.bracket != 'main':
            continue
        sources = meta.get('sources') or []
        if len(sources) != 2:
            continue
        kinds = {src.get('kind') for src in sources if isinstance(src, dict)}
        if not ({'winner', 'loser'} & kinds):
            candidates.append(m)
    if not candidates:
        return []
    min_depth = min(int(meta_for(m).get('depth', 0)) for m in candidates)
    return [m for m in candidates if int(meta_for(m).get('depth', 0)) == min_depth]



def player_group_map(t):
    """Return player_id -> group key based on Group Stage fixtures."""
    out = {}
    for m in Match.query.filter_by(tournament_id=t.id).all():
        if not m.bracket.startswith('group_'):
            continue
        if m.player1_id:
            out[m.player1_id] = m.bracket
        if m.player2_id:
            out[m.player2_id] = m.bracket
    return out


def tournament_seed_rows(t, player_ids=None):
    """Stable post-group tournament ranking used for reseeding.

    Ranking is based only on Group Stage performance so the seed remains
    auditable throughout the knockout: points/match, win rate, set diff/match,
    point diff/match, then original seed.
    """
    rows = []
    group_map = player_group_map(t)
    allowed = set(player_ids) if player_ids is not None else None
    for group_key, table in compute_group_tables(t):
        for pos, row in enumerate(table, start=1):
            p = row['player']
            if allowed is not None and p.id not in allowed:
                continue
            metrics = comparison_metrics(row)
            matches = max(1, metrics['matches'])
            rows.append({
                'player': p,
                'group': group_map.get(p.id, group_key),
                'group_position': pos,
                'ppm': metrics['points'] / matches,
                'win_rate': metrics['wins'] / matches,
                'set_diff_pm': metrics['set_diff'] / matches,
                'point_diff_pm': metrics['point_diff'] / matches,
            })
    # Tournament seed: group winners are ranked ahead of runners-up, who
    # are ranked ahead of qualified third-placed teams. Inside each tier use
    # normalized group-stage performance, then original seed as the audit-safe
    # final tiebreaker.
    rows.sort(key=lambda r: (r['group_position'], -r['ppm'], -r['win_rate'], -r['set_diff_pm'], -r['point_diff_pm'], r['player'].seed))
    for idx, row in enumerate(rows, start=1):
        row['tournament_seed'] = idx
    return rows


def recommended_pairs(t, player_ids):
    """Reseed strongest vs weakest while avoiding Group Stage rematches.

    For an 8-team field this creates a top-half vs bottom-half draw. The
    strongest surviving team is matched with the weakest compatible opponent,
    then the next strongest, etc. A small backtracking search finds a no-rematch
    solution when one exists. If impossible, it minimizes rematches and keeps
    the strongest-vs-weakest preference.
    """
    ranked = tournament_seed_rows(t, player_ids)
    players = [r['player'] for r in ranked]
    if len(players) < 2 or len(players) % 2:
        return []
    groups = player_group_map(t)
    half = len(players) // 2
    top = players[:half]
    bottom = list(reversed(players[half:]))  # weakest first

    best = None
    # fields used here are typically <= 16; permutations of the bottom half
    # are still small enough for this organizer workflow.
    if len(bottom) <= 8:
        permutations = itertools.permutations(bottom)
    else:
        permutations = [tuple(bottom)]
    for perm in permutations:
        rematches = sum(1 for a, b in zip(top, perm) if groups.get(a.id) and groups.get(a.id) == groups.get(b.id))
        # preference penalty: keep weak opponents aligned with strong seeds
        rank_index = {p.id: i for i, p in enumerate(bottom)}
        preference = sum(abs(i - rank_index[b.id]) for i, b in enumerate(perm))
        score = (rematches, preference)
        if best is None or score < best[0]:
            best = (score, list(zip(top, perm)))
            if score == (0, 0):
                break
    return best[1] if best else []


def source_match_keys(meta):
    keys = []
    for src in meta.get('sources') or []:
        if isinstance(src, dict) and src.get('kind') in ('winner', 'loser') and src.get('match_key'):
            keys.append(src['match_key'])
    return keys


def downstream_match_ids(t, origin_matches):
    """Return descendant match ids fed by any of origin_matches."""
    matches = Match.query.filter_by(tournament_id=t.id).all()
    key_to_match = {meta_for(m).get('key'): m for m in matches if meta_for(m).get('key')}
    origins = {meta_for(m).get('key') for m in origin_matches if meta_for(m).get('key')}
    descendants = set()
    frontier = set(origins)
    changed = True
    while changed:
        changed = False
        for m in matches:
            if m.id in descendants:
                continue
            deps = set(source_match_keys(meta_for(m)))
            if deps & frontier:
                descendants.add(m.id)
                k = meta_for(m).get('key')
                if k:
                    frontier.add(k)
                changed = True
    return descendants


def editable_round_groups(t):
    """Ready post-entry rounds that can be manually reseeded.

    The first Championship round has its dedicated editor. Later-round editing
    is exposed only when Settings -> Pairing mode is "Reseed every round".
    """
    if settings_for(t).get('pairing_mode', 'first_round_only') != 'every_round':
        return []
    if t.system == 'group_knockout' and not group_stage_complete(t.id):
        return []

    entry_ids = {m.id for m in championship_entry_matches(t)} if t.system == 'group_knockout' else set()
    buckets = defaultdict(list)
    matches = Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index, Match.id).all()
    for m in matches:
        if m.bracket.startswith('group_') or m.id in entry_ids:
            continue
        meta = meta_for(m)
        if not m.player1_id or not m.player2_id:
            continue
        section = meta.get('section') or m.round_name
        depth = int(meta.get('depth', 0))
        buckets[(section, depth)].append(m)

    seed_lookup = {r['player'].id: r for r in tournament_seed_rows(t)}
    out = []
    for (section, depth), group in buckets.items():
        group = sorted(group, key=lambda m: (meta_for(m).get('slot_index', 0), m.order_index, m.id))
        if len(group) < 2:
            continue
        unique_ids = []
        for m in group:
            unique_ids.extend([m.player1_id, m.player2_id])
        if len(set(unique_ids)) != len(unique_ids):
            continue
        locked = any(meta_for(m).get('user_result_entered') is True for m in group)
        descendants = downstream_match_ids(t, group)
        downstream_started = any(
            meta_for(dm).get('user_result_entered') is True
            for dm in (Match.query.get(mid) for mid in descendants) if dm
        )
        recommended = recommended_pairs(t, unique_ids)
        out.append({
            'key': f'{section}|{depth}',
            'section': section,
            'depth': depth,
            'matches': group,
            'players': [Player.query.get(pid) for pid in unique_ids],
            'seed_rows': [seed_lookup.get(pid) for pid in unique_ids if seed_lookup.get(pid)],
            'recommended': recommended,
            'locked': locked or downstream_started,
        })
    out.sort(key=lambda x: (min(m.order_index for m in x['matches']), x['section'], x['depth']))
    return out

def apply_smart_byes_after_groups(t):
    """Assign unavoidable BYEs to the highest post-group Tournament Seeds.

    Example: 10 qualifiers require a 16-slot logical draw. Only four teams
    need to play a preliminary round, while Tournament Seeds #1-#6 receive a
    BYE into the 8-team core bracket. The existing exact-placement tree is
    preserved; only the entrant sources are rewritten after the Group Stage.
    """
    settings = settings_for(t)
    if settings.get('bye_policy', 'top_seeds') != 'top_seeds':
        return False
    qualifiers = championship_qualifier_rows(t)
    n = len(qualifiers)
    if n < 3 or n & (n - 1) == 0:
        return False
    ranked = tournament_seed_rows(t, [q['player'].id for q in qualifiers])
    players = [r['player'] for r in ranked]
    core_size = 1 << (n.bit_length() - 1)
    play_in_matches = n - core_size
    bye_count = n - 2 * play_in_matches
    if play_in_matches <= 0 or bye_count < 0:
        return False

    all_matches = Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index, Match.id).all()
    prelim = []
    for m in all_matches:
        meta = meta_for(m)
        if m.bracket == 'main' and meta.get('section') == 'Championship' and ('playin' in str(meta.get('key','')).lower() or 'preliminary' in (m.round_name or '').lower()):
            prelim.append(m)
    if len(prelim) != play_in_matches:
        return False

    # Strongest seeds are protected. Remaining entrants are paired strongest
    # vs weakest, with the existing no-group-rematch optimizer where possible.
    bye_players = players[:bye_count]
    play_players = players[bye_count:]
    pairs = recommended_pairs(t, [p.id for p in play_players])
    if len(pairs) != play_in_matches:
        # deterministic fallback
        pairs=[]
        for i in range(play_in_matches):
            pairs.append((play_players[i], play_players[-1-i]))
    for m, (p1,p2) in zip(prelim, pairs):
        meta=meta_for(m)
        if meta.get('manual_pairing'):
            continue
        meta['sources']=[{'kind':'player','player_id':p1.id},{'kind':'player','player_id':p2.id}]
        meta['protected_bye_draw']=True
        set_meta(m,meta)
        m.player1_id=p1.id; m.player2_id=p2.id
        m.score_json='[]'; m.winner_id=None; m.loser_id=None

    prelim_keys={meta_for(m).get('key') for m in prelim}
    # The first core round contains a mixture of direct entrants (BYEs) and
    # winners of preliminary matches. Replace only the direct entrant slots.
    core=[]
    for m in all_matches:
        meta=meta_for(m)
        if m.bracket!='main' or meta.get('section')!='Championship':
            continue
        srcs=meta.get('sources') or []
        if any(isinstance(src,dict) and src.get('kind')=='winner' and src.get('match_key') in prelim_keys for src in srcs):
            core.append(m)
    direct_slots=[]
    for m in core:
        meta=meta_for(m); srcs=list(meta.get('sources') or [])
        for idx,src in enumerate(srcs):
            if not (isinstance(src,dict) and src.get('kind')=='winner' and src.get('match_key') in prelim_keys):
                direct_slots.append((m,idx))
    if len(direct_slots) != len(bye_players):
        return False
    for (m,idx),p in zip(direct_slots, bye_players):
        meta=meta_for(m); srcs=list(meta.get('sources') or [])
        srcs[idx]={'kind':'player','player_id':p.id,'protected_bye':True,'tournament_seed':next((r['tournament_seed'] for r in ranked if r['player'].id==p.id),None)}
        meta['sources']=srcs; meta['protected_bye_draw']=True; set_meta(m,meta)

    db.session.commit()
    propagate_all(t.id)
    assign_schedule(t, renumber=False)
    return True


def apply_recommended_entry_pairing(t):
    """Apply recommended first knockout draw once Group Stage is complete."""
    if settings_for(t).get('pairing_mode', 'first_round_only') == 'fixed_bracket':
        return False
    qualifiers = championship_qualifier_rows(t)
    entry = championship_entry_matches(t)
    if not qualifiers or len(entry) * 2 != len(qualifiers):
        return False
    ids = [q['player'].id for q in qualifiers]
    pairs = recommended_pairs(t, ids)
    if len(pairs) != len(entry):
        return False
    for m, (p1, p2) in zip(entry, pairs):
        meta = meta_for(m)
        # Respect an organizer's already-saved draw.
        if meta.get('manual_pairing'):
            return False
        meta['sources'] = [{'kind': 'player', 'player_id': p1.id}, {'kind': 'player', 'player_id': p2.id}]
        meta['auto_reseeded'] = True
        set_meta(m, meta)
        m.player1_id = p1.id
        m.player2_id = p2.id
        m.score_json = '[]'
        m.winner_id = None
        m.loser_id = None
    db.session.commit()
    propagate_all(t.id)
    assign_schedule(t, renumber=False)
    return True


def score_sets(score_json):
    try:
        return json.loads(score_json or '[]')
    except Exception:
        return []

def generate_single_elimination(t: Tournament):
    clear_matches(t)
    phase = next(p for p in t.phases if p.phase_type == 'main')
    order_state = {'value': 0}
    slots = [{'kind': 'player', 'player_id': p.id} if p else {'kind': 'empty'} for p in seeded_bracket(t.players)]
    if t.consolation:
        build_exact_bracket(t.id, phase.id, slots, 1, 'overall', order_state, bracket='main')
    else:
        # classic championship only
        current = slots
        depth = 0
        while len(current) > 1:
            nxt = []
            section = 'Championship'
            rn = round_label(section, depth, int(math.log2(len(current))))
            for i in range(0, len(current), 2):
                key = f'main_d{depth}_m{i//2+1}'
                order_state['value'] += 1
                add_match(t.id, phase.id, order_state['value'], 'main', rn, current[i], current[i+1], section, 1, len(current), depth, i // 2, key)
                nxt.append({'kind': 'winner', 'match_key': key})
            current = nxt
            depth += 1
    db.session.commit()
    propagate_all(t.id)
    assign_schedule(t, renumber=True)


def generate_group_knockout(t: Tournament):
    clear_matches(t)
    players = sorted(t.players, key=lambda p: p.seed)
    group_phase = next(p for p in t.phases if p.phase_type == 'groups')
    knockout_phase = next(p for p in t.phases if p.phase_type == 'main')
    s = settings_for(t)
    qualifiers_per_group = max(1, int(s.get('qualifiers_per_group', 2)))
    best_next_place_qualifiers = max(0, int(s.get('best_next_place_qualifiers', 0)))
    group_count = max(1, min(int(s.get('group_count', 2)), len(players)))
    assignment_mode = s.get('group_assignment_mode', 'auto')
    groups = [[] for _ in range(group_count)]

    if assignment_mode == 'manual':
        manual = s.get('manual_groups') or {}
        unassigned = []
        for p in players:
            try:
                group_no = int(manual.get(str(p.id), 0))
            except Exception:
                group_no = 0
            if 1 <= group_no <= group_count:
                groups[group_no - 1].append(p)
            else:
                unassigned.append(p)
        if unassigned:
            names = ', '.join(p.name for p in unassigned[:8])
            if len(unassigned) > 8:
                names += ', ...'
            raise ValueError(f'Manual group assignment is incomplete. Assign a group to: {names}')
    else:
        idx = 0
        direction = 1
        while idx < len(players):
            order = list(range(group_count)) if direction == 1 else list(range(group_count - 1, -1, -1))
            for g in order:
                if idx < len(players):
                    groups[g].append(players[idx])
                    idx += 1
            direction *= -1

    for g in groups:
        g.sort(key=lambda p: p.seed)
    if any(len(g) < 2 for g in groups):
        bad = [chr(65+i) for i, g in enumerate(groups) if len(g) < 2]
        raise ValueError('Each group must contain at least 2 players. Check Group ' + ', '.join(bad) + '.')
    if qualifiers_per_group > min(len(g) for g in groups):
        raise ValueError('Qualifiers per group cannot exceed the size of the smallest group.')
    if best_next_place_qualifiers > group_count:
        raise ValueError('Best next-place qualifiers cannot exceed the number of groups.')
    next_position = qualifiers_per_group + 1
    groups_with_next_position = sum(1 for g in groups if len(g) >= next_position)
    if best_next_place_qualifiers > groups_with_next_position:
        raise ValueError(
            f'Only {groups_with_next_position} groups contain a team in position {next_position}. '
            f'You cannot qualify {best_next_place_qualifiers} teams from that position.'
        )
    championship_field = group_count * qualifiers_per_group + best_next_place_qualifiers
    # The number of GROUPS does not need to be a power of two. What matters is
    # the size of each elimination/placement field. build_exact_bracket_any()
    # handles non-power-of-two fields with preliminary matches and byes.
    # Example: 13 teams / 3 groups / top 2 + two best 3rd-place teams =
    # 8-team Championship plus a 5-team Places 9-13 bracket.

    order_idx = 0
    for gi, group_players in enumerate(groups, start=1):
        for rnd, pairs in enumerate(round_robin_pairings(group_players), start=1):
            for a, b in pairs:
                order_idx += 1
                db.session.add(Match(
                    tournament_id=t.id,
                    phase_id=group_phase.id,
                    round_name=f'Group Round {rnd}',
                    bracket=f'group_{gi}',
                    order_index=order_idx,
                    player1_id=a.id,
                    player2_id=b.id,
                ))
    db.session.commit()

    group_tables = compute_group_tables(t)
    order_state = {'value': order_idx}

    if qualifiers_per_group >= 2:
        # Keep the established cross-group championship seeding for the first two positions.
        champ_sources = build_group_sources(group_tables, [1, 2])
        if qualifiers_per_group > 2:
            for pos_extra in range(3, qualifiers_per_group + 1):
                champ_sources.extend(
                    {'kind': 'group_pos', 'group': f'group_{i+1}', 'pos': pos_extra}
                    for i in range(len(group_tables))
                )
    else:
        champ_sources = [{'kind': 'group_pos', 'group': f'group_{i+1}', 'pos': 1} for i in range(len(group_tables))]

    # Optional wildcards: e.g. 2/group qualify + two best 3rd-placed teams.
    wildcard_position = qualifiers_per_group + 1
    for wildcard_rank in range(1, best_next_place_qualifiers + 1):
        champ_sources.append({'kind': 'best_group_pos', 'pos': wildcard_position, 'rank': wildcard_rank})

    champ_size = len(champ_sources)
    if champ_size >= 2:
        if t.consolation:
            build_exact_bracket_any(t.id, knockout_phase.id, champ_sources, 1, 'groups_champ', order_state, bracket='main')
        else:
            build_championship_bracket_any(t.id, knockout_phase.id, champ_sources, 'groups_champ', order_state, bracket='main')

    # Exact placement for EVERY non-qualified team. Keep all remaining teams
    # in one placement field instead of splitting them into artificial 2-place
    # bands. This is important with unequal group sizes (e.g. 13 teams in
    # 3 groups -> 5/4/4): the five non-qualifiers form Places 9-13.
    # build_exact_bracket_any() then creates the required preliminary match /
    # byes and the loser of that preliminary match can finish 13th on court.
    if t.consolation:
        placement_sources = []
        max_group_size = max(len(g) for g in groups)
        for pos in range(qualifiers_per_group + 1, max_group_size + 1):
            candidates_at_pos = sum(1 for _key, rows in group_tables if len(rows) >= pos)
            skip = best_next_place_qualifiers if pos == wildcard_position else 0
            for rank in range(skip + 1, candidates_at_pos + 1):
                placement_sources.append({'kind': 'best_group_pos', 'pos': pos, 'rank': rank})

        if len(placement_sources) >= 2:
            next_place = champ_size + 1
            build_exact_bracket_any(
                t.id, knockout_phase.id, placement_sources, next_place,
                f'place_{next_place}', order_state, bracket='placement'
            )

    db.session.commit()
    propagate_all(t.id)
    assign_schedule(t, renumber=True)


def generate_double_elimination(t: Tournament):
    # kept intentionally simple; main focus in this build is exact consolation / placement logic
    generate_single_elimination(t)


def generate_round_robin(t: Tournament):
    clear_matches(t)
    phase = t.phases[0]
    order_idx = 0
    for rnd, pairs in enumerate(round_robin_pairings(sorted(t.players, key=lambda p: p.seed)), start=1):
        for a, b in pairs:
            order_idx += 1
            db.session.add(Match(tournament_id=t.id, phase_id=phase.id, round_name=f'Round {rnd}', bracket='league', order_index=order_idx, player1_id=a.id, player2_id=b.id))
    db.session.commit()
    assign_schedule(t, renumber=True)


def generate_swiss(t: Tournament):
    clear_matches(t)
    phase = t.phases[0]
    players = sorted(t.players, key=lambda p: p.seed)
    rounds = max(3, math.ceil(math.log2(max(2, len(players)))))
    order_idx = 0
    for rnd in range(1, rounds + 1):
        for i in range(0, len(players), 2):
            p1 = players[i]
            p2 = players[i + 1] if i + 1 < len(players) else None
            order_idx += 1
            db.session.add(Match(tournament_id=t.id, phase_id=phase.id, round_name=f'Swiss Round {rnd}', bracket='swiss', order_index=order_idx, player1_id=p1.id if p1 else None, player2_id=p2.id if p2 else None))
        players = players[1:] + players[:1]
    db.session.commit()
    assign_schedule(t, renumber=True)



def expected_phase_types(system):
    return {
        'group_knockout': ['groups', 'main'],
        'single_elimination': ['main'],
        'double_elimination': ['main', 'losers', 'final'],
        'round_robin': ['groups'],
        'swiss': ['swiss', 'placement'],
    }.get(system, ['main'])


def ensure_phases_for_system(t: Tournament):
    """Repair phase definitions if a tournament was created by an older build.

    Generation buttons should always work even when an existing database came
    from a previous version of the application. Phase rules are preserved when
    the current phase layout is already compatible.
    """
    current_types = [p.phase_type for p in sorted(t.phases, key=lambda p: p.order_index)]
    expected = expected_phase_types(t.system)
    if current_types == expected:
        return
    # Existing matches can point at obsolete phases, so clear them first.
    Match.query.filter_by(tournament_id=t.id).delete()
    db.session.flush()
    t.phases.clear()
    db.session.flush()
    create_default_phases(t)
    db.session.commit()


def sync_player_count(t: Tournament, requested_count: int):
    """Keep the Players setting and the actual player list in sync.

    Older builds only changed Tournament.player_count, which could leave the
    Generate Structure action working on a different number of players than
    the UI displayed.
    """
    requested_count = max(2, int(requested_count))
    players = sorted(list(t.players), key=lambda p: p.seed)
    if requested_count > len(players):
        for seed in range(len(players) + 1, requested_count + 1):
            db.session.add(Player(tournament_id=t.id, name=f'Player {seed}', seed=seed, club=''))
    elif requested_count < len(players):
        keep = players[:requested_count]
        remove = players[requested_count:]
        for p in remove:
            db.session.delete(p)
        current = settings_for(t)
        manual = current.get('manual_groups') or {}
        keep_ids = {str(p.id) for p in keep}
        current['manual_groups'] = {k: v for k, v in manual.items() if k in keep_ids}
        save_settings_dict(t, current)
    t.player_count = requested_count
    db.session.flush()


def generate_structure(t):
    ensure_phases_for_system(t)
    if t.system == 'group_knockout':
        generate_group_knockout(t)
    elif t.system == 'single_elimination':
        generate_single_elimination(t)
    elif t.system == 'double_elimination':
        generate_double_elimination(t)
    elif t.system == 'round_robin':
        generate_round_robin(t)
    else:
        generate_swiss(t)


def group_stage_complete(tournament_id):
    """Return True only when every scheduled group-stage match has a winner.

    Post-group brackets must never be populated from provisional standings.
    This is especially important when wildcard qualifiers (e.g. best 3rd-place
    teams) are used: until the final group match is completed, every
    Championship/Placement slot sourced from a group stays unresolved.
    """
    group_matches = (
        Match.query
        .filter(Match.tournament_id == tournament_id, Match.bracket.like('group_%'))
        .all()
    )
    return bool(group_matches) and all(m.winner_id is not None for m in group_matches)


def resolve_source(source, tournament_id, matches_by_key, group_tables_by_key, groups_complete=True):
    if not source:
        return None
    kind = source.get('kind')
    if kind == 'empty':
        return None
    if kind == 'player':
        return source.get('player_id')
    if kind == 'winner':
        m = matches_by_key.get(source.get('match_key'))
        return m.winner_id if m else None
    if kind == 'loser':
        m = matches_by_key.get(source.get('match_key'))
        return m.loser_id if m else None
    if kind == 'group_pos':
        if not groups_complete:
            return None
        rows = group_tables_by_key.get(source.get('group'), [])
        pos = int(source.get('pos', 0))
        if 1 <= pos <= len(rows):
            return rows[pos - 1]['player'].id
    if kind == 'best_group_pos':
        if not groups_complete:
            return None
        pos = int(source.get('pos', 0))
        rank = int(source.get('rank', 0))
        candidates = best_group_position_rows(list(group_tables_by_key.items()), pos)
        if 1 <= rank <= len(candidates):
            return candidates[rank - 1][1]['player'].id
    return None


def source_state(source, tournament_id, matches_by_key, group_tables_by_key,
                 groups_complete=True, visiting=None):
    """Return (player_id, resolved). A pending feeder is never an empty slot."""
    if not source or source.get('kind') == 'empty':
        return None, True
    kind = source.get('kind')
    if kind in ('player', 'group_pos', 'best_group_pos'):
        if kind != 'player' and not groups_complete:
            return None, False
        return resolve_source(source, tournament_id, matches_by_key,
                              group_tables_by_key, groups_complete), True
    if kind not in ('winner', 'loser'):
        return None, False
    key = source.get('match_key')
    visiting = set(visiting or ())
    if key in visiting or key not in matches_by_key:
        return None, False
    visiting.add(key)
    feeder = matches_by_key[key]
    inputs = meta_for(feeder).get('sources') or []
    if inputs:
        states = [source_state(src, tournament_id, matches_by_key,
                               group_tables_by_key, groups_complete, visiting)
                  for src in inputs]
        if not all(ready for _, ready in states):
            return None, False
        participants = [pid for pid, _ in states if pid is not None]
        if len(participants) < 2:
            # A proven structural BYE has a winner but no loser.
            return (participants[0] if participants and kind == 'winner' else None), True
        if [feeder.player1_id, feeder.player2_id] != [pid for pid, _ in states]:
            return None, False
    if not feeder.winner_id:
        return None, False
    return (feeder.winner_id if kind == 'winner' else feeder.loser_id), True


def propagate_all(tournament_id):
    t = Tournament.query.get(tournament_id)
    if not t:
        return
    matches = Match.query.filter_by(tournament_id=tournament_id).order_by(Match.order_index, Match.id).all()
    matches_by_key = {}
    for m in matches:
        meta = meta_for(m)
        key = meta.get('key')
        if key:
            matches_by_key[key] = m

    loops = 0
    while loops < 20:
        loops += 1
        changed = False
        group_tables_by_key = {k: rows for k, rows in compute_group_tables(t)}
        groups_complete = group_stage_complete(tournament_id) if t.system == 'group_knockout' else True
        for m in matches:
            meta = meta_for(m)
            sources = meta.get('sources') or []
            if sources:
                states = [source_state(src, tournament_id, matches_by_key,
                                       group_tables_by_key, groups_complete) for src in sources]
                new_p1 = states[0][0] if states else None
                new_p2 = states[1][0] if len(states) > 1 else None
                if m.player1_id != new_p1 or m.player2_id != new_p2:
                    m.player1_id = new_p1
                    m.player2_id = new_p2
                    m.score_json = '[]'
                    m.winner_id = None
                    m.loser_id = None
                    changed = True

            inputs_ready = not sources or all(ready for _, ready in states)
            # Clear legacy automatic wins while an upstream outcome is pending.
            if not inputs_ready and (m.winner_id or m.loser_id):
                m.winner_id = None
                m.loser_id = None
                changed = True
            if inputs_ready and m.player1_id and not m.player2_id:
                if m.winner_id != m.player1_id or m.loser_id is not None:
                    m.winner_id = m.player1_id
                    m.loser_id = None
                    changed = True
            elif inputs_ready and m.player2_id and not m.player1_id:
                if m.winner_id != m.player2_id or m.loser_id is not None:
                    m.winner_id = m.player2_id
                    m.loser_id = None
                    changed = True
            elif not m.player1_id and not m.player2_id:
                if m.winner_id or m.loser_id or m.score_json != '[]':
                    m.winner_id = None
                    m.loser_id = None
                    m.score_json = '[]'
                    changed = True
        if not changed:
            break
    db.session.commit()


def repair_unplayed_brackets(tournament_id=None):
    repaired = []
    query = Tournament.query.filter_by(id=tournament_id) if tournament_id is not None else Tournament.query
    for t in query.all():
        matches = list(t.matches)
        generated = [m for m in matches if meta_for(m).get('sources')]
        if not generated or any(score_sets(m.score_json) or meta_for(m).get('user_result_entered') for m in generated):
            continue
        # Idempotent: only reconsider tournaments with legacy unscored wins.
        if not any(m.winner_id for m in generated):
            continue
        before = [(m.id, m.player1_id, m.player2_id, m.winner_id, m.loser_id) for m in generated]
        propagate_all(t.id)
        after = [(m.id, m.player1_id, m.player2_id, m.winner_id, m.loser_id) for m in generated]
        if before != after:
            repaired.append(t.id)
    return repaired


def closing_match_kind(match):
    meta = meta_for(match)
    band = (meta.get('place_start'), meta.get('place_end'))
    if band == (1, 2):
        return 'final'
    if band == (3, 4):
        return 'bronze'
    return None


def interleave_knockout_order(matches, court_count):
    """Fill parallel court slots with main/placement matches whose feeders are earlier."""
    groups = [m for m in matches if m.bracket.startswith('group_')]
    remaining = [m for m in matches if not m.bracket.startswith('group_')]
    if not any(meta_for(m).get('key') for m in remaining):
        return matches
    rank = {m.id: i for i, m in enumerate(matches)}
    keyed = {meta_for(m).get('key'): m.id for m in matches if meta_for(m).get('key')}
    dependencies = {m.id: {keyed[src['match_key']] for src in meta_for(m).get('sources', [])
                           if src.get('kind') in ('winner', 'loser') and src.get('match_key') in keyed}
                    for m in remaining}
    done = {m.id for m in groups}
    ordered = list(groups)
    lane_cursor = 0
    while remaining:
        ordinary = [m for m in remaining if not closing_match_kind(m)]
        pool = ordinary or [m for m in remaining if closing_match_kind(m) == 'bronze'] or remaining
        ready = sorted([m for m in pool if dependencies[m.id] <= done], key=lambda m: rank[m.id])
        if not ready:
            raise ValueError('Cannot schedule bracket: missing or cyclic match dependency.')
        batch, occupied = [], set()
        for _ in range(court_count):
            eligible = [m for m in ready if m not in batch and not ({m.player1_id,m.player2_id} - {None}) & occupied]
            if not eligible:
                break
            preferred = [m for m in eligible if (0 if meta_for(m).get('place_start') == 1 else 1) == lane_cursor]
            chosen = (preferred or eligible)[0]
            batch.append(chosen)
            occupied.update({chosen.player1_id,chosen.player2_id} - {None})
            lane_cursor = 1 - (0 if meta_for(chosen).get('place_start') == 1 else 1)
            if closing_match_kind(chosen):
                break # Bronze and final each get a separate closing slot.
        ordered.extend(batch)
        done.update(m.id for m in batch)
        remaining = [m for m in remaining if m not in batch]
    return ordered


def allocate_court_slots(matches, court_count):
    wave, court = 1, 0
    keys_in_wave, players_in_wave = set(), set()
    previous_closing = False
    for m in matches:
        meta = meta_for(m)
        deps = {src.get('match_key') for src in meta.get('sources', []) if src.get('kind') in ('winner','loser')}
        players = {m.player1_id,m.player2_id} - {None}
        closing = closing_match_kind(m)
        if court and (court >= court_count or deps & keys_in_wave or players & players_in_wave or closing or previous_closing):
            wave += 1
            court = 0
            keys_in_wave, players_in_wave = set(), set()
        court += 1
        m.court_number = court
        meta['schedule_wave'] = wave
        set_meta(m,meta)
        if meta.get('key'):
            keys_in_wave.add(meta['key'])
        players_in_wave.update(players)
        previous_closing = bool(closing)


def upgrade_unplayed_schedules():
    for t in Tournament.query.all():
        if settings_for(t).get('schedule_version') == 2:
            continue
        knockout = [m for m in t.matches if meta_for(m).get('key')]
        if not knockout or any(score_sets(m.score_json) or meta_for(m).get('user_result_entered') for m in knockout):
            continue
        assign_schedule(t,renumber=True)


def assign_schedule(tournament, renumber=False):
    """Assign courts while keeping published match numbers stable.

    During structure generation ``renumber=True`` performs one dependency-safe
    numbering pass. After the tournament starts, calls use ``renumber=False``:
    pairings can change and participants can propagate, but match #16 never
    suddenly becomes match #8.
    """
    matches = Match.query.filter_by(tournament_id=tournament.id).order_by(Match.order_index, Match.id).all()
    court_count = max(1, tournament.court_count)
    if not matches:
        return

    if renumber:
        group_round_max = 0
        for m in matches:
            if m.bracket.startswith('group_'):
                try:
                    group_round_max = max(group_round_max, int((m.round_name or '').split()[-1]))
                except Exception:
                    group_round_max = max(group_round_max, 1)

        key_to_match = {}
        for m in matches:
            key = meta_for(m).get('key')
            if key:
                key_to_match[key] = m

        level_cache = {}
        def match_level(match):
            if match.id in level_cache:
                return level_cache[match.id]
            if match.bracket.startswith('group_'):
                try:
                    level = int((match.round_name or '').split()[-1]) - 1
                except Exception:
                    level = 0
                level_cache[match.id] = max(0, level)
                return level_cache[match.id]
            sources = meta_for(match).get('sources') or []
            deps = []
            for src in sources:
                kind = (src or {}).get('kind')
                if kind in ('player', 'empty'):
                    deps.append(-1)
                elif kind in ('group_pos', 'best_group_pos'):
                    deps.append(group_round_max - 1 if group_round_max > 0 else 0)
                elif kind in ('winner', 'loser'):
                    dep = key_to_match.get((src or {}).get('match_key'))
                    deps.append(match_level(dep) if dep else 0)
                else:
                    deps.append(0)
            level = max(deps) + 1 if deps else 0
            level_cache[match.id] = level
            return level

        grouped = defaultdict(list)
        for m in matches:
            grouped[match_level(m)].append(m)
        ordered = []
        for level in sorted(grouped):
            level_matches = grouped[level]
            group_matches = [m for m in level_matches if m.bracket.startswith('group_')]
            other_matches = [m for m in level_matches if not m.bracket.startswith('group_')]
            if group_matches:
                by_group = defaultdict(list)
                for m in sorted(group_matches, key=lambda x: (x.bracket, x.id)):
                    by_group[m.bracket].append(m)
                buckets = [by_group[k] for k in sorted(by_group)]
                for i in range(max(len(b) for b in buckets)):
                    for bucket in buckets:
                        if i < len(bucket):
                            ordered.append(bucket[i])
            def other_key(m):
                meta = meta_for(m)
                section = meta.get('section') or ''
                start = meta.get('place_start', 999)
                slot = meta.get('slot_index', 0)
                section_rank = 0 if section == 'Championship' else (1 if section == '3rd Place' else 2)
                return (section_rank, start, slot, m.id)
            ordered.extend(sorted(other_matches, key=other_key))
        ordered = interleave_knockout_order(ordered, court_count)
        for idx, m in enumerate(ordered, start=1):
            m.order_index = idx
        matches = sorted(ordered, key=lambda m: m.order_index)

    # Keep published numbers on normal score saves, while separating dependencies
    # and the two closing matches into different parallel court slots.
    allocate_court_slots(sorted(matches, key=lambda m: (m.order_index,m.id)), court_count)
    if renumber:
        settings = settings_for(tournament)
        settings['schedule_version'] = 2
        save_settings_dict(tournament,settings)
    db.session.commit()

def placement_map(t):
    """Return exact places that can be proven from completed placement paths.

    In addition to ordinary 2-team placement finals (e.g. 3rd/4th, 5th/6th),
    odd-sized exact-placement brackets may start with a preliminary match.
    If exactly one loser from a placement band has no downstream loser path,
    that player owns the last place in that band (e.g. loser of the 9-11
    preliminary is 11th). This keeps Final Standings exact after every match
    has been completed.
    """
    placements = {}
    matches = Match.query.filter_by(tournament_id=t.id).all()

    # Track which match outcomes are consumed by later matches.
    winner_consumed = set()
    loser_consumed = set()
    keyed = {}
    for m in matches:
        meta = meta_for(m)
        key = meta.get('key')
        if key:
            keyed[key] = m
        for src in meta.get('sources', []) or []:
            mk = src.get('match_key')
            if not mk:
                continue
            if src.get('kind') == 'winner':
                winner_consumed.add(mk)
            elif src.get('kind') == 'loser':
                loser_consumed.add(mk)

    # Standard exact placement finals.
    for m in matches:
        meta = meta_for(m)
        start_place = meta.get('place_start')
        end_place = meta.get('place_end')
        if start_place is None or end_place is None:
            continue
        if end_place - start_place == 1 and m.winner_id:
            placements[m.winner_id] = start_place
            if m.loser_id:
                placements[m.loser_id] = end_place

    # Odd field sizes: a single terminal preliminary loser can own the final
    # place in a wider placement band, such as Places 9-11 or 9-13.
    bands = defaultdict(list)
    for m in matches:
        meta = meta_for(m)
        start_place = meta.get('place_start')
        end_place = meta.get('place_end')
        if start_place is None or end_place is None or end_place - start_place <= 1:
            continue
        bands[(start_place, end_place)].append(m)

    for (start_place, end_place), band_matches in bands.items():
        terminal_losers = []
        for m in band_matches:
            meta = meta_for(m)
            key = meta.get('key')
            if m.loser_id and key and key not in loser_consumed:
                terminal_losers.append(m)
        # Only infer an exact last place when there is a single unambiguous
        # terminal loser. Multiple terminal losers would represent a range and
        # must be resolved by additional placement matches.
        if len(terminal_losers) == 1:
            placements.setdefault(terminal_losers[0].loser_id, end_place)

    # fallback for league / swiss
    if t.system in ('round_robin', 'swiss'):
        table = standings_for_matches(matches, sorted(t.players, key=lambda p: p.seed))
        return {row['player'].id: idx for idx, row in enumerate(table, start=1)}
    return placements


def tournament_is_complete(t):
    """True when every playable match in the generated structure is resolved.

    BYEs count as resolved when the engine has assigned a winner. A future
    match with unresolved participants keeps the tournament provisional.
    """
    matches = Match.query.filter_by(tournament_id=t.id).all()
    if not matches:
        return False
    for m in matches:
        meta = meta_for(m)
        sources = meta.get('sources', []) or []
        if m.player1_id and m.player2_id:
            if not m.winner_id:
                return False
        elif m.player1_id or m.player2_id:
            # A one-sided generated match is a BYE only when already resolved.
            if not m.winner_id:
                return False
        elif sources:
            # The structure still contains a future match whose feeder results
            # have not resolved into participants yet.
            return False
    return True


def final_ranking(t):
    players = sorted(t.players, key=lambda p: p.seed)
    places = placement_map(t)
    group_rank = {}
    for _, rows in compute_group_tables(t):
        for i, row in enumerate(rows, start=1):
            group_rank[row['player'].id] = i
    return sorted(players, key=lambda p: (places.get(p.id, 9999), group_rank.get(p.id, 9999), p.seed))


def export_database_json(t):
    data = {
        'tournament': {
            'name': t.name, 'player_count': t.player_count, 'system': t.system,
            'consolation': t.consolation, 'court_count': t.court_count, 'settings_json': t.settings_json,
        },
        'players': [{'name': p.name, 'club': p.club, 'seed': p.seed} for p in sorted(t.players, key=lambda x: x.seed)],
        'phases': [{'name': ph.name, 'phase_type': ph.phase_type, 'order_index': ph.order_index, 'best_of': ph.best_of, 'points_to': ph.points_to, 'config_json': ph.config_json} for ph in t.phases],
        'matches': [{
            'phase_order': m.phase.order_index if m.phase else None,
            'round_name': m.round_name, 'bracket': m.bracket, 'order_index': m.order_index, 'court_number': m.court_number,
            'player1_seed': m.player1.seed if m.player1 else None,
            'player2_seed': m.player2.seed if m.player2 else None,
            'score_json': m.score_json, 'winner_seed': m.winner.seed if m.winner else None, 'loser_seed': m.loser.seed if m.loser else None,
            'notes': m.notes
        } for m in Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index).all()]
    }
    return json.dumps(data, ensure_ascii=False, indent=2)


def import_database_json(payload):
    raw = json.loads(payload)
    td = raw['tournament']
    t = Tournament(name=td['name'], player_count=td['player_count'], system=td['system'], consolation=td['consolation'], court_count=td['court_count'], settings_json=td.get('settings_json', '{}'))
    imported_settings = settings_for(t)
    imported_settings.pop('public_token', None)
    imported_settings.pop('public_slug', None)
    imported_settings['public_enabled'] = False
    save_settings_dict(t, imported_settings)
    db.session.add(t)
    db.session.flush()
    for p in raw.get('players', []):
        db.session.add(Player(tournament_id=t.id, name=p['name'], club=p.get('club', ''), seed=p['seed']))
    db.session.flush()
    seed_to_player = {p.seed: p for p in Player.query.filter_by(tournament_id=t.id).all()}
    phases = []
    for ph in raw.get('phases', []):
        phase = Phase(tournament_id=t.id, **ph)
        db.session.add(phase)
        phases.append(phase)
    db.session.flush()
    by_order = {ph.order_index: ph for ph in phases}
    old_id_to_new = {}
    for saved in raw.get('matches', []):
        metadata = json.loads(saved.get('notes') or '{}')
        for slot, source in enumerate(metadata.get('sources') or []):
            seed = saved.get(f'player{slot + 1}_seed')
            if source.get('kind') == 'player' and seed in seed_to_player:
                old_id_to_new[source.get('player_id')] = seed_to_player[seed].id
    for m in raw.get('matches', []):
        metadata = json.loads(m.get('notes') or '{}')
        for source in metadata.get('sources') or []:
            if source.get('kind') == 'player':
                source['player_id'] = old_id_to_new.get(source.get('player_id'))
        db.session.add(Match(
            tournament_id=t.id,
            phase_id=by_order.get(m.get('phase_order')).id if by_order.get(m.get('phase_order')) else None,
            round_name=m['round_name'], bracket=m['bracket'], order_index=m['order_index'], court_number=m['court_number'],
            player1_id=seed_to_player[m['player1_seed']].id if m.get('player1_seed') in seed_to_player else None,
            player2_id=seed_to_player[m['player2_seed']].id if m.get('player2_seed') in seed_to_player else None,
            score_json=m.get('score_json', '[]'),
            winner_id=seed_to_player[m['winner_seed']].id if m.get('winner_seed') in seed_to_player else None,
            loser_id=seed_to_player[m['loser_seed']].id if m.get('loser_seed') in seed_to_player else None,
            notes=json.dumps(metadata, ensure_ascii=False),
        ))
    db.session.commit()
    repair_unplayed_brackets(t.id)
    return t


def bracket_sections(t):
    """Return draw-sheet friendly visual sections.

    Knockout matches are split into logical draws instead of one mixed graph:
    Championship, 3rd Place and exact-placement bands (5-8, 9-16, ...).
    Every draw still keeps source metadata, so an empty first slot can say
    ``Loser #16`` / ``Winner #20`` before the participant is known.
    """
    matches = Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index, Match.id).all()
    group_sections = defaultdict(lambda: defaultdict(list))
    draw_buckets = defaultdict(list)

    def placement_band(place_start):
        try:
            p = int(place_start)
        except Exception:
            return None
        if p <= 2:
            return ('championship', 1)
        if p <= 4:
            return ('bronze', 3)
        # Exact-placement families are 5-8, 9-16, 17-32, ...
        lo = (1 << int(math.floor(math.log2(max(1, p - 1))))) + 1
        return ('placement', lo)

    for m in matches:
        meta = meta_for(m)
        if m.bracket.startswith('group_'):
            name = m.bracket.replace('group_', 'Group ')
            try:
                depth = int((m.round_name or '').split()[-1]) - 1
            except Exception:
                depth = 0
            group_sections[name][depth].append(m)
            continue

        section = str(meta.get('section') or '')
        ps = meta.get('place_start')
        if section == 'Championship' or ps in (1, 2):
            key = ('championship', 1)
        elif section == '3rd Place' or ps in (3, 4):
            key = ('bronze', 3)
        else:
            key = placement_band(ps) or ('other', 999)
        draw_buckets[key].append(m)

    ordered = []
    if group_sections:
        sections=[]
        for name in sorted(group_sections, key=lambda x: int(x.split()[-1]) if x.split()[-1].isdigit() else x):
            rounds=[sorted(group_sections[name][d], key=lambda m:(m.order_index,m.id)) for d in sorted(group_sections[name])]
            sections.append({'name': name, 'rounds': rounds, 'kind': 'group'})
        ordered.append({'stage':'Group Stage','sections':sections})

    if draw_buckets:
        sections=[]
        key_order = sorted(draw_buckets, key=lambda k: (0 if k[0]=='championship' else 1 if k[0]=='bronze' else 2, k[1]))
        for key in key_order:
            bucket = draw_buckets[key]
            depths = sorted({int(meta_for(m).get('depth', 0) or 0) for m in bucket})
            rounds=[]
            for d in depths:
                rm = [m for m in bucket if int(meta_for(m).get('depth',0) or 0)==d]
                rounds.append(sorted(rm, key=lambda m:(meta_for(m).get('slot_index',0), m.order_index, m.id)))
            if key[0]=='championship':
                name='Championship Draw'
            elif key[0]=='bronze':
                name='3rd Place'
            else:
                starts=[int(meta_for(m).get('place_start',999) or 999) for m in bucket]
                ends=[int(meta_for(m).get('place_end',0) or 0) for m in bucket]
                lo=min(starts) if starts else key[1]
                hi=max(ends) if ends else lo
                name=label_for_places(lo, hi)
            sections.append({'name':name,'rounds':rounds,'kind':key[0]})
        ordered.append({'stage':'Tournament Draws','sections':sections})
    return ordered

def source_label(match, slot_index):
    meta = meta_for(match)
    sources = meta.get('sources') or []
    if slot_index >= len(sources):
        return 'TBD'
    src = sources[slot_index] or {}
    kind = src.get('kind')
    if kind == 'group_pos':
        g = str(src.get('group','group_?')).replace('group_','Group ')
        return f"{g} #{src.get('pos','?')}"
    if kind == 'best_group_pos':
        return f"Best #{src.get('pos','?')} · rank {src.get('rank','?')}"
    if kind in ('winner','loser'):
        key = src.get('match_key')
        dep = None
        if key:
            for m in Match.query.filter_by(tournament_id=match.tournament_id).all():
                if meta_for(m).get('key') == key:
                    dep = m; break
        prefix = 'Winner' if kind == 'winner' else 'Loser'
        return f"{prefix} #{dep.order_index}" if dep else f"{prefix} previous match"
    if kind == 'player':
        p = Player.query.get(src.get('player_id'))
        return p.name if p else 'TBD'
    if kind == 'empty':
        return 'BYE'
    return 'TBD'


def logo_path_for(t):
    filename = settings_for(t).get('logo_filename','')
    if not filename:
        return None
    path = os.path.join(UPLOAD_DIR, os.path.basename(filename))
    return path if os.path.isfile(path) else None


_PDF_UNICODE_FONT = None


def _pdf_unicode_font():
    """Return a system TrueType font with Polish glyph support.

    The application intentionally does not bundle font files. On Windows we
    prefer Arial/Segoe UI/Calibri; on Linux/macOS we fall back to common system
    fonts. This keeps the PDF portable while rendering 'Michał' as a real glyph.
    """
    global _PDF_UNICODE_FONT
    if _PDF_UNICODE_FONT:
        return _PDF_UNICODE_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [
        ('CreatorArial', r'C:\Windows\Fonts\arial.ttf'),
        ('CreatorSegoe', r'C:\Windows\Fonts\segoeui.ttf'),
        ('CreatorCalibri', r'C:\Windows\Fonts\calibri.ttf'),
        ('CreatorDejaVu', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
        ('CreatorLiberation', '/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'),
        ('CreatorArialMac', '/Library/Fonts/Arial.ttf'),
    ]
    for name, path in candidates:
        if os.path.isfile(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                _PDF_UNICODE_FONT = name
                return name
            except Exception:
                continue
    # Last-resort fallback. Modern Windows installs should always hit Arial or Segoe.
    _PDF_UNICODE_FONT = 'Helvetica'
    return _PDF_UNICODE_FONT


def _draw_creator_credit(c, width, y=4.8*mm, size=6.8):
    """Draw the fixed creator credit right-aligned in the PDF footer."""
    font = _pdf_unicode_font()
    c.setFont(font, size)
    c.drawRightString(width - 8 * mm, y, CREATOR_CREDIT)


def draw_pdf_logo(c, t, width, height):
    """Draw the fixed creator credit and optional tournament logo without overlap."""
    c.saveState()
    c.setFillGray(0.55)
    _draw_creator_credit(c, width, 4.8 * mm, 6.8)
    c.restoreState()

    path = logo_path_for(t)
    if not path:
        return
    try:
        img = ImageReader(path)
        iw, ih = img.getSize()
        max_w, max_h = 18*mm, 10*mm
        scale = min(max_w/iw, max_h/ih)
        w, h = iw*scale, ih*scale
        # Keep the logo above the creator line so both remain readable.
        c.drawImage(img, width-8*mm-w, 8*mm, width=w, height=h, preserveAspectRatio=True, mask='auto')
    except Exception:
        pass

def _pdf_title(c, title, subtitle, width, height):
    c.setFont('Helvetica-Bold', 18)
    c.drawString(14 * mm, height - 14 * mm, title[:90])
    c.setFont('Helvetica', 10)
    c.setFillGray(0.35)
    c.drawString(14 * mm, height - 21 * mm, subtitle[:120])
    c.setFillGray(0)
    c.setLineWidth(0.4)
    c.line(14 * mm, height - 25 * mm, width - 14 * mm, height - 25 * mm)


def _draw_round_title(round_matches, round_index, round_count, draw_name):
    if not round_matches:
        return f'Round {round_index + 1}'
    if draw_name == 'Championship Draw':
        remaining = round_count - round_index
        names = {1: 'Final', 2: 'Semi Finals', 3: 'Quarter Finals', 4: 'Round of 16', 5: 'Round of 32'}
        return names.get(remaining, f'Round {round_index + 1}')
    rn = round_matches[0].round_name or f'Round {round_index + 1}'
    if ' - ' in rn:
        rn = rn.split(' - ', 1)[-1]
    return rn[:30]


def _section_height_mm(section):
    max_count = max((len(r) for r in section.get('rounds', [])), default=1)
    # Taller cards + larger type. A draw can move to a new A4 page rather than shrink.
    return max(34, max_count * 21 + 18)


def _resolve_pdf_round_centers(raw_items, top_y, bottom_y, box_h, gap_y):
    """Place match boxes with guaranteed non-overlap while staying near dependency centers."""
    if not raw_items:
        return []
    items=sorted(raw_items, key=lambda x:x['desired'], reverse=True)
    centers=[]
    prev_bottom=None
    for item in items:
        center=min(top_y-box_h/2, max(bottom_y+box_h/2, item['desired']))
        if prev_bottom is not None:
            center=min(center, prev_bottom-box_h-gap_y)
        centers.append([item,center])
        prev_bottom=center
    low=min(c for _,c in centers)-box_h/2
    if low < bottom_y:
        shift=bottom_y-low
        for pair in centers: pair[1]+=shift
    high=max(c for _,c in centers)+box_h/2
    if high > top_y:
        shift=high-top_y
        for pair in centers: pair[1]-=shift
    # Guaranteed final spacing.
    for i in range(1,len(centers)):
        max_center=centers[i-1][1]-box_h-gap_y
        if centers[i][1]>max_center: centers[i][1]=max_center
    return centers


def _render_pdf_draw(c, t, section, x0, y_top, width_avail, height_avail):
    rounds = section.get('rounds') or []
    if not rounds:
        return
    c.setFont('Helvetica-Bold', 11)
    c.drawString(x0, y_top, section['name'][:70])
    top = y_top - 7 * mm
    bottom = y_top - height_avail + 3 * mm
    round_count = len(rounds)
    gap = 8 * mm
    col_w = min(55 * mm, (width_avail-gap*max(0,round_count-1))/max(1,round_count))
    if round_count>1:
        gap=max(7*mm,(width_avail-col_w*round_count)/(round_count-1))
    box_h=14*mm; gap_y=3*mm
    keypos={}; layouts=[]
    for ri,rmatches in enumerate(rounds):
        x=x0+ri*(col_w+gap)
        title=_draw_round_title(rmatches,ri,round_count,section['name'])
        c.setFont('Helvetica-Bold',8.1); c.setFillGray(.34); c.drawCentredString(x+col_w/2,top+2.8*mm,title); c.setFillGray(0)
        usable_top=top-3*mm; usable_bottom=bottom+1*mm
        raw=[]
        for mi,m in enumerate(rmatches):
            if ri==0:
                usable=max(1,usable_top-usable_bottom)
                desired=usable_top-usable*((mi+.5)/max(1,len(rmatches)))
            else:
                deps=[]
                for src in (meta_for(m).get('sources') or []):
                    if (src or {}).get('kind') in ('winner','loser'):
                        kp=keypos.get((src or {}).get('match_key'))
                        if kp: deps.append(kp[1])
                if deps: desired=sum(deps)/len(deps)
                else:
                    usable=max(1,usable_top-usable_bottom)
                    desired=usable_top-usable*((mi+.5)/max(1,len(rmatches)))
            raw.append({'m':m,'desired':desired})
        resolved=_resolve_pdf_round_centers(raw,usable_top,usable_bottom,box_h,gap_y)
        for item,y in resolved:
            m=item['m']; k=meta_for(m).get('key')
            if k:keypos[k]=(x,y,col_w)
            layouts.append((ri,m,x,y,col_w))

    c.setStrokeGray(.65); c.setLineWidth(.55)
    for ri,m,tx,ty,cw in layouts:
        if ri==0: continue
        for src in (meta_for(m).get('sources') or []):
            if (src or {}).get('kind') not in ('winner','loser'): continue
            sp=keypos.get((src or {}).get('match_key'))
            if not sp: continue
            sx,sy,sw=sp; sx+=sw
            if tx<=sx: continue
            bend=sx+(tx-sx)*.5
            c.line(sx,sy,bend,sy); c.line(bend,sy,bend,ty); c.line(bend,ty,tx,ty)
    c.setStrokeGray(0)

    for ri,m,x,y,cw in layouts:
        card_y=y-box_h/2
        c.setFillGray(1); c.setStrokeGray(.62); c.setLineWidth(.45); c.roundRect(x,card_y,cw,box_h,1.4*mm,stroke=1,fill=1)
        c.setStrokeGray(.8); c.line(x,y,x+cw,y); c.setStrokeGray(0); c.setFillGray(0)
        p1=m.player1.name if m.player1 else source_label(m,0); p2=m.player2.name if m.player2 else source_label(m,1)
        c.setFont('Helvetica-Bold' if m.winner_id==m.player1_id and m.player1_id else 'Helvetica',8.2); c.drawString(x+1.7*mm,y+1.7*mm,p1[:34])
        c.setFont('Helvetica-Bold' if m.winner_id==m.player2_id and m.player2_id else 'Helvetica',8.2); c.drawString(x+1.7*mm,y-5.0*mm,p2[:34])
        c.setFont('Helvetica',6.3); c.setFillGray(.42); meta_text=f'#{m.order_index}'; score=score_to_text(m.score_json)
        if score: meta_text+=f'  {score}'
        c.drawRightString(x+cw-1.4*mm,card_y+1.5*mm,meta_text[:42]); c.setFillGray(0)


def write_bracket_pdf(t):
    """A4 landscape draw sheets. Readability wins over squeezing all paths onto one page."""
    buf=io.BytesIO(); pagesize=landscape(A4); c=canvas.Canvas(buf,pagesize=pagesize)
    width,height=pagesize; c.setTitle(f'{t.name} - Tournament Draws')
    stages=[x for x in bracket_sections(t) if x['stage']!='Group Stage']; sections=stages[0]['sections'] if stages else []
    if not sections:
        _pdf_title(c,t.name,'Tournament Draws',width,height); c.setFont('Helvetica',11); c.drawString(14*mm,height-40*mm,'No knockout bracket has been generated yet.'); draw_pdf_logo(c,t,width,height); c.save(); buf.seek(0); return buf
    margin_x=10*mm; usable_w=width-20*mm; y=height-31*mm; page_open=False
    for section in sections:
        band_h=_section_height_mm(section)*mm
        max_band=height-48*mm
        band_h=min(band_h,max_band)
        if (not page_open) or y-band_h<14*mm:
            if page_open:
                draw_pdf_logo(c,t,width,height); c.showPage()
            _pdf_title(c,t.name,'Tournament Draws - Championship + exact placement',width,height)
            y=height-31*mm; page_open=True
        _render_pdf_draw(c,t,section,margin_x,y,usable_w,band_h)
        y-=band_h+5*mm
    c.setFont('Helvetica',6.5); c.setFillGray(.45); c.drawString(10*mm,6.5*mm,'Future slots show Winner / Loser of match #.'); c.setFillGray(0); draw_pdf_logo(c,t,width,height)
    c.save(); buf.seek(0); return buf


def _short_stage_name(stage):
    st=(stage or '').lower()
    if 'group' in st:return 'Group Stage'
    if 'championship' in st or 'placement' in st:return 'Knockout / Placement'
    return stage or ''


def write_groups_pdf(t):
    """A4 landscape group sheets: standing table left, match list right, up to 4 groups/page."""
    buf=io.BytesIO(); pagesize=landscape(A4); c=canvas.Canvas(buf,pagesize=pagesize)
    width,height=pagesize; c.setTitle(f'{t.name} - Groups & Results'); tables=compute_group_tables(t)
    if not tables:
        _pdf_title(c,t.name,'Groups & Results',width,height); c.setFont('Helvetica',10); c.drawString(12*mm,height-36*mm,'No group tables available.'); draw_pdf_logo(c,t,width,height); c.save(); buf.seek(0); return buf
    mx=9*mm; top=height-29*mm; bottom=9*mm; gap=4*mm; groups_per_page=4
    block_h=(top-bottom-gap*(groups_per_page-1))/groups_per_page
    def head(): _pdf_title(c,t.name,'Group Stage - Tables & Match Results',width,height)
    head()
    for gi,(gname,rows) in enumerate(tables):
        slot=gi%groups_per_page
        if gi and slot==0:
            draw_pdf_logo(c,t,width,height); c.showPage(); head()
        ytop=top-slot*(block_h+gap); ybot=ytop-block_h
        title=gname.replace('group_','Group ')
        c.setFillGray(.965); c.roundRect(mx, ybot, width-2*mx, block_h, 1.7*mm, stroke=0, fill=1); c.setFillGray(0)
        c.setFont('Helvetica-Bold',10.5); c.drawString(mx+3*mm,ytop-5*mm,title)
        split=width*.57; sx=mx+3*mm; sy=ytop-11*mm
        headers=['#','Team','M','W','L','Pts','Sets','PD']; rel=[0,7,61,73,84,95,108,126]
        c.setFont('Helvetica-Bold',7.8)
        for h,dx in zip(headers,rel): c.drawString(sx+dx*mm,sy,h)
        c.setStrokeGray(.75); c.line(sx,sy-2*mm,split-4*mm,sy-2*mm); c.setStrokeGray(0)
        y=sy-6*mm; c.setFont('Helvetica',8.5)
        for idx,r in enumerate(rows,1):
            vals=[idx,r['player'].name,r['matches'],r['wins'],r['losses'],r['points'],f"{r['sets_won']}:{r['sets_lost']}",r['point_diff']]
            limits=[2,28,2,2,2,3,6,6]
            for v,dx,lim in zip(vals,rel,limits): c.drawString(sx+dx*mm,y,str(v)[:lim])
            y-=4.2*mm
        # Match list on the right; no grid lines through text.
        rx=split+2*mm; ry=ytop-5*mm; c.setFont('Helvetica-Bold',8.8); c.drawString(rx,ry,'Matches'); ry-=5*mm
        gm=Match.query.filter_by(tournament_id=t.id,bracket=gname).order_by(Match.order_index).all(); c.setFont('Helvetica',8.0)
        for m in gm:
            p1=m.player1.name if m.player1 else 'TBD'; p2=m.player2.name if m.player2 else 'TBD'; sc=score_to_text(m.score_json) or '-'
            c.drawString(rx,ry,f"#{m.order_index}  {p1}  {sc}  {p2}"[:72]); ry-=4.1*mm
    draw_pdf_logo(c,t,width,height); c.save(); buf.seek(0); return buf


def write_schedule_pdf(t):
    """A4 landscape schedule with section bands and readable 8.5pt rows."""
    matches=Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index,Match.id).all()
    pagesize=landscape(A4); buf=io.BytesIO(); c=canvas.Canvas(buf,pagesize=pagesize); width,height=pagesize; c.setTitle(f'{t.name} - Match Schedule')
    mx=8*mm; top=height-28*mm; bottom=8*mm; row_h=5.0*mm; section_h=5.5*mm
    cols=[('No.',0,5),('Court',12,5),('Round',28,38),('Team 1',98,28),('Team 2',154,28),('Score',210,43)]
    page=1
    def header():
        _pdf_title(c,t.name,'Match Schedule',width,height); c.setFont('Helvetica',7); c.setFillGray(.45); c.drawRightString(width-mx,height-18*mm,f'Page {page}'); c.setFillGray(0)
        y=top; c.setFont('Helvetica-Bold',8.3)
        for txt,x,_ in cols:c.drawString(mx+x*mm,y,txt)
        c.setStrokeGray(.55); c.line(mx,y-2*mm,width-mx,y-2*mm); c.setStrokeGray(0); return y-6*mm
    y=header(); prev_stage=None
    for m in matches:
        stage=_short_stage_name(m.phase.name if m.phase else '')
        need= row_h + (section_h if stage!=prev_stage else 0)
        if y-need<bottom:
            draw_pdf_logo(c,t,width,height); c.showPage(); page+=1; y=header(); prev_stage=None
        if stage!=prev_stage:
            c.setFillGray(.94); c.roundRect(mx,y-section_h+1.1*mm,width-2*mx,section_h-1.4*mm,1*mm,stroke=0,fill=1); c.setFillGray(0)
            c.setFont('Helvetica-Bold',8.4); c.drawString(mx+2*mm,y-2.3*mm,stage.upper()); y-=section_h
        vals=[m.order_index,m.court_number,m.round_name,m.player1.name if m.player1 else 'TBD',m.player2.name if m.player2 else 'TBD',score_to_text(m.score_json) or '']
        c.setFont('Helvetica',8.5)
        for v,(_,x,lim) in zip(vals,cols): c.drawString(mx+x*mm,y,str(v)[:lim])
        c.setStrokeGray(.88); c.setLineWidth(.25); c.line(mx,y-1.8*mm,width-mx,y-1.8*mm); c.setStrokeGray(0); y-=row_h; prev_stage=stage
    draw_pdf_logo(c,t,width,height); c.save(); buf.seek(0); return buf


def write_final_standings_pdf(t):
    """Printable A4 final standings. Uses exact placement results when available."""
    buf = io.BytesIO()
    pagesize = A4
    c = canvas.Canvas(buf, pagesize=pagesize)
    width, height = pagesize
    c.setTitle(f'{t.name} - Final Standings')

    ranking = final_ranking(t)
    places = placement_map(t)
    # FINAL means the generated tournament structure has been fully played.
    # It must not depend on whether every place happened to be inferred by an
    # older placement-map implementation.
    exact_known = tournament_is_complete(t)
    mx = 16 * mm
    top = height - 31 * mm
    bottom = 13 * mm
    row_h = 8.0 * mm
    page = 1

    def header():
        _pdf_title(c, t.name, 'Final Standings', width, height)
        c.setFont('Helvetica', 7.2)
        c.setFillGray(.48)
        status = 'FINAL' if exact_known else 'CURRENT / PROVISIONAL'
        c.drawRightString(width - mx, height - 20.5 * mm, f'{status} · Page {page}')
        c.setFillGray(0)
        y = top
        c.setFillGray(.94)
        c.roundRect(mx, y - 7 * mm, width - 2 * mx, 7 * mm, 1 * mm, stroke=0, fill=1)
        c.setFillGray(0)
        c.setFont('Helvetica-Bold', 9.2)
        c.drawString(mx + 2 * mm, y - 4.7 * mm, 'Place')
        c.drawString(mx + 25 * mm, y - 4.7 * mm, 'Team')
        c.drawString(mx + 105 * mm, y - 4.7 * mm, 'Club')
        c.drawRightString(width - mx - 2 * mm, y - 4.7 * mm, 'Original seed')
        return y - 11 * mm

    y = header()
    for idx, p in enumerate(ranking, start=1):
        if y - row_h < bottom:
            draw_pdf_logo(c, t, width, height)
            c.showPage()
            page += 1
            y = header()
        place = places.get(p.id)
        place_text = str(place if place is not None else idx)
        c.setFont('Helvetica-Bold', 10.2)
        c.drawString(mx + 2 * mm, y, place_text)
        c.setFont('Helvetica-Bold', 10.2 if place and place <= 3 else 9.8)
        c.drawString(mx + 25 * mm, y, (p.name or '')[:42])
        c.setFont('Helvetica', 8.8)
        c.setFillGray(.25)
        c.drawString(mx + 105 * mm, y, (p.club or '')[:30])
        c.drawRightString(width - mx - 2 * mm, y, f'#{p.seed}')
        c.setFillGray(0)
        c.setStrokeGray(.86)
        c.setLineWidth(.3)
        c.line(mx, y - 2.8 * mm, width - mx, y - 2.8 * mm)
        c.setStrokeGray(0)
        y -= row_h

    if not ranking:
        c.setFont('Helvetica', 10)
        c.drawString(mx, y, 'No standings available yet.')

    draw_pdf_logo(c, t, width, height)
    c.save()
    buf.seek(0)
    return buf


def write_results_excel(t):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = 'Standings'
    ws1.append(['Place', 'Player', 'Club', 'Seed'])
    ranking = final_ranking(t)
    places = placement_map(t)
    for i, p in enumerate(ranking, start=1):
        ws1.append([places.get(p.id, i), p.name, p.club, p.seed])
    ws2 = wb.create_sheet('Matches')
    ws2.append(['Order', 'Court', 'Phase', 'Round', 'Bracket', 'Player 1', 'Player 2', 'Score', 'Winner'])
    for m in Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index).all():
        ws2.append([m.order_index, m.court_number, m.phase.name if m.phase else '', m.round_name, m.bracket,
                    m.player1.name if m.player1 else '', m.player2.name if m.player2 else '', score_to_text(m.score_json), m.winner.name if m.winner else ''])
    if t.system == 'group_knockout':
        for group_name, rows in compute_group_tables(t):
            ws = wb.create_sheet(group_name.replace('group_', 'Group '))
            ws.append(['Place', 'Player', 'Matches', 'Wins', 'Losses', 'Points', 'Sets +', 'Sets -', 'Set Diff', 'Points +', 'Points -'])
            for i, row in enumerate(rows, start=1):
                p = row['player']
                ws.append([i, p.name, row['matches'], row['wins'], row['losses'], row['points'], row['sets_won'], row['sets_lost'], row['set_diff'], row['points_won'], row['points_lost']])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0); return buf


def import_players_from_excel(tournament_id, file_storage):
    wb = load_workbook(file_storage, data_only=True)
    ws = wb.active
    headers = {}
    for idx, cell in enumerate(ws[1], start=1):
        value = str(cell.value).strip() if cell.value is not None else ''
        headers[value.lower()] = idx
    rank_col = headers.get('rank') or headers.get('ranking') or headers.get('seed')
    name_col = headers.get('name') or headers.get('player') or headers.get('competitor')
    if not rank_col or not name_col:
        raise ValueError('The file must contain Rank and Name columns.')
    rows = []
    for r in range(2, ws.max_row + 1):
        rank = ws.cell(r, rank_col).value
        name = ws.cell(r, name_col).value
        if rank is None or name is None:
            continue
        rows.append((int(rank), str(name).strip()))
    rows.sort(key=lambda x: x[0])
    t = Tournament.query.get_or_404(tournament_id)
    Player.query.filter_by(tournament_id=t.id).delete()
    t.player_count = len(rows)
    current = settings_for(t)
    current['manual_groups'] = {}
    save_settings_dict(t, current)
    db.session.flush()
    for idx, (_, name) in enumerate(rows, start=1):
        db.session.add(Player(tournament_id=t.id, name=name, seed=idx, club=''))
    db.session.commit()


@app.route('/help')
def help_page():
    return render_template('help.html')


@app.route('/')
def index():
    tournaments = Tournament.query.order_by(Tournament.id.desc()).all()
    return render_template('index.html', tournaments=tournaments, systems=POPULAR_SYSTEMS)


@app.route('/create', methods=['POST'])
def create():
    name = request.form.get('name', 'New Tournament')
    player_count = int(request.form.get('player_count', 8))
    system = request.form.get('system', 'single_elimination')
    consolation = request.form.get('consolation') == 'on'
    court_count = int(request.form.get('court_count', 1))
    t = create_tournament(name, player_count, system, consolation, court_count)
    return redirect(url_for('tournament_view', tournament_id=t.id))


@app.route('/tournament/<int:tournament_id>')
def tournament_view(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    matches = Match.query.filter_by(tournament_id=t.id).order_by(Match.order_index).all()
    ranking = final_ranking(t) if matches else []
    group_tables = compute_group_tables(t)
    sections = bracket_sections(t)
    settings = settings_for(t)
    group_count = max(1, int(settings.get('group_count', 2)))
    group_names = [chr(65 + i) if i < 26 else str(i + 1) for i in range(group_count)]
    championship_field = group_count * max(1, int(settings.get('qualifiers_per_group', 2))) + max(0, int(settings.get('best_next_place_qualifiers', 0)))
    placement_field = max(0, len(t.players) - championship_field)
    groups_complete = group_stage_complete(t.id) if t.system == 'group_knockout' and matches else False
    knockout_qualifiers = championship_qualifier_rows(t) if groups_complete else []
    knockout_entry_matches = championship_entry_matches(t) if groups_complete else []
    manual_pairing_ready = settings.get('pairing_mode', 'first_round_only') != 'fixed_bracket' and bool(knockout_qualifiers) and len(knockout_entry_matches) * 2 == len(knockout_qualifiers)
    round_pairings = editable_round_groups(t)
    tournament_seeds = tournament_seed_rows(t, [q['player'].id for q in knockout_qualifiers]) if knockout_qualifiers else []
    return render_template('tournament.html' if g.user and g.user.role == 'admin' and not getattr(g, 'public_view', False) else 'readonly.html', tournament=t, systems=POPULAR_SYSTEMS, matches=matches, ranking=ranking, group_tables=group_tables, bracket_sections=sections, score_to_text=score_to_text, score_sets=score_sets, meta_for=meta_for, source_match_keys=source_match_keys, source_label=source_label, settings=settings, placement_map=placement_map(t), group_names=group_names, championship_field=championship_field, placement_field=placement_field, groups_complete=groups_complete, knockout_qualifiers=knockout_qualifiers, knockout_entry_matches=knockout_entry_matches, manual_pairing_ready=manual_pairing_ready, round_pairings=round_pairings, tournament_seeds=tournament_seeds, logo_url=(url_for('tournament_logo', tournament_id=t.id) if logo_path_for(t) else None), draw_round_title=_draw_round_title, comparison_locked=any(score_sets(m.score_json) for m in matches))


@app.route('/tournament/<int:tournament_id>/settings', methods=['POST'])
def update_settings(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    current = settings_for(t)
    comparison = request.form.get('wildcard_comparison', current.get('wildcard_comparison', 'ppm'))
    if comparison not in ('ppm', 'exclude_bottom'):
        abort(400)
    if comparison != current.get('wildcard_comparison', 'ppm') and any(score_sets(m.score_json) for m in t.matches):
        flash('Metodę porównywania grup wybierz przed wpisaniem pierwszego wyniku. Trwający turniej zachowuje dotychczasowe zasady.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#settings')
    current['wildcard_comparison'] = comparison
    current['qualifiers_per_group'] = max(1, int(request.form.get('qualifiers_per_group', current.get('qualifiers_per_group', 2))))
    current['best_next_place_qualifiers'] = max(0, int(request.form.get('best_next_place_qualifiers', current.get('best_next_place_qualifiers', 0))))
    current['group_count'] = max(1, int(request.form.get('group_count', current.get('group_count', 2))))
    mode = request.form.get('group_assignment_mode', current.get('group_assignment_mode', 'auto'))
    current['group_assignment_mode'] = mode if mode in ('auto', 'manual') else 'auto'
    pairing_mode = request.form.get('pairing_mode', current.get('pairing_mode', 'first_round_only'))
    current['pairing_mode'] = pairing_mode if pairing_mode in ('first_round_only', 'every_round', 'fixed_bracket') else 'first_round_only'
    current['bye_policy'] = 'top_seeds' if request.form.get('bye_policy', current.get('bye_policy', 'top_seeds')) == 'top_seeds' else 'structure_order'
    current.setdefault('manual_groups', {})
    save_settings_dict(t, current)
    t.name = request.form.get('name', t.name)
    old_system = t.system
    t.system = request.form.get('system', t.system)
    requested_count = int(request.form.get('player_count', t.player_count))
    t.court_count = max(1, int(request.form.get('court_count', t.court_count)))
    t.consolation = request.form.get('consolation') == 'on'
    sync_player_count(t, requested_count)
    db.session.commit()
    if old_system != t.system:
        ensure_phases_for_system(t)
    if request.form.get('generate_after') == '1':
        try:
            generate_structure(t)
            flash('Settings saved and tournament structure generated successfully.')
        except Exception as e:
            db.session.rollback()
            flash(f'Could not generate structure: {e}')
    else:
        flash('Settings saved.')
    return redirect(url_for('tournament_view', tournament_id=t.id))


@app.route('/tournament/<int:tournament_id>/players', methods=['POST'])
def update_players(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    payload = request.get_json(force=True)
    order = payload.get('players', [])
    by_id = {p.id: p for p in t.players}
    for idx, item in enumerate(order, start=1):
        p = by_id.get(int(item['id']))
        if not p:
            continue
        p.seed = idx
        p.name = item.get('name', p.name)
        p.club = item.get('club', p.club)
    db.session.commit()
    return jsonify({'ok': True})


@app.route('/tournament/<int:tournament_id>/groups', methods=['POST'])
def update_group_assignments(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    payload = request.get_json(force=True)
    current = settings_for(t)
    group_count = max(1, int(payload.get('group_count', current.get('group_count', 2))))
    mode = payload.get('group_assignment_mode', current.get('group_assignment_mode', 'manual'))
    assignments = payload.get('manual_groups') or {}
    valid_ids = {str(p.id) for p in t.players}
    cleaned = {}
    for pid, group_no in assignments.items():
        if str(pid) not in valid_ids:
            continue
        try:
            group_no = int(group_no)
        except Exception:
            continue
        if 1 <= group_no <= group_count:
            cleaned[str(pid)] = group_no
    current['group_count'] = group_count
    current['group_assignment_mode'] = mode if mode in ('auto', 'manual') else 'manual'
    current['manual_groups'] = cleaned
    save_settings_dict(t, current)
    db.session.commit()
    return jsonify({'ok': True})


@app.route('/tournament/<int:tournament_id>/players/import_excel', methods=['POST'])
def import_players_excel(tournament_id):
    f = request.files.get('file')
    if not f:
        flash('Choose an Excel file.')
        return redirect(url_for('tournament_view', tournament_id=tournament_id))
    try:
        import_players_from_excel(tournament_id, f)
        flash('Ranking imported from Excel.')
    except Exception as e:
        flash(f'Excel import error: {e}')
    return redirect(url_for('tournament_view', tournament_id=tournament_id))


@app.route('/tournament/<int:tournament_id>/generate', methods=['POST'])
def generate(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    try:
        generate_structure(t)
        flash('Tournament structure generated successfully.')
    except Exception as e:
        db.session.rollback()
        flash(f'Could not generate structure: {e}')
    return redirect(url_for('tournament_view', tournament_id=t.id))


@app.route('/tournament/<int:tournament_id>/knockout_pairing', methods=['POST'])
def update_knockout_pairing(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    if t.system != 'group_knockout' or not group_stage_complete(t.id):
        flash('Complete all Group Stage matches before editing knockout pairings.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#knockout-pairing')
    if settings_for(t).get('pairing_mode', 'first_round_only') == 'fixed_bracket':
        flash('Manual first-round pairing is disabled because Fixed bracket mode is selected.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#knockout-pairing')

    qualifiers = championship_qualifier_rows(t)
    qualifier_ids = {item['player'].id for item in qualifiers}
    entry_matches = championship_entry_matches(t)
    if not qualifiers or len(entry_matches) * 2 != len(qualifiers):
        flash('Manual pairing is available when the Championship starts with one complete knockout round. This structure uses play-ins/byes and needs a different slot editor.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#knockout-pairing')

    selected = []
    pairs = []
    try:
        for m in entry_matches:
            p1 = int(request.form.get(f'match_{m.id}_p1', '0'))
            p2 = int(request.form.get(f'match_{m.id}_p2', '0'))
            pairs.append((m, p1, p2))
            selected.extend([p1, p2])
    except Exception:
        flash('Invalid knockout pairing selection.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#knockout-pairing')

    if set(selected) != qualifier_ids or len(selected) != len(qualifier_ids) or len(set(selected)) != len(selected):
        flash('Each qualified team must be used exactly once in the first knockout round.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#knockout-pairing')

    # Do not allow a Group Stage rematch in the first knockout round when a
    # zero-rematch draw is mathematically available.
    groups = player_group_map(t)
    selected_rematches = sum(1 for _m, p1, p2 in pairs if groups.get(p1) and groups.get(p1) == groups.get(p2))
    recommended = recommended_pairs(t, selected)
    minimum_rematches = sum(1 for a, b in recommended if groups.get(a.id) and groups.get(a.id) == groups.get(b.id)) if recommended else selected_rematches
    if selected_rematches > minimum_rematches:
        flash('This first-round draw contains an avoidable rematch between teams from the same Group Stage group.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#knockout-pairing')

    # Manual pairings remain editable until the user has ACTUALLY submitted
    # a post-group result. Never infer this from winner_id, score_json, BYEs,
    # or propagated slots: those can be generated automatically by the engine.
    post_group_matches = Match.query.filter(
        Match.tournament_id == t.id,
        ~Match.bracket.like('group_%')
    ).all()
    knockout_played = None
    for candidate in post_group_matches:
        cmeta = meta_for(candidate)
        if cmeta.get('user_result_entered') is True:
            knockout_played = candidate
            break
    if knockout_played:
        flash(f'Knockout pairings are locked because Match #{knockout_played.order_index} already has a user-entered result. Clear that result before changing the draw.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#knockout-pairing')

    # Remove every automatically propagated post-group state before applying a
    # revised draw. This prevents stale winners/BYEs from making a second save
    # behave as if the knockout had already started. Group-stage data is untouched.
    entry_ids = {m.id for m, _, _ in pairs}
    for candidate in post_group_matches:
        cmeta = meta_for(candidate)
        if candidate.id not in entry_ids:
            candidate.player1_id = None
            candidate.player2_id = None
        candidate.score_json = '[]'
        candidate.winner_id = None
        candidate.loser_id = None
        cmeta.pop('user_result_entered', None)
        set_meta(candidate, cmeta)

    for m, p1, p2 in pairs:
        meta = meta_for(m)
        meta['sources'] = [
            {'kind': 'player', 'player_id': p1},
            {'kind': 'player', 'player_id': p2},
        ]
        meta['manual_pairing'] = True
        set_meta(m, meta)
        m.player1_id = p1
        m.player2_id = p2
        m.score_json = '[]'
        m.winner_id = None
        m.loser_id = None

    db.session.commit()
    propagate_all(t.id)
    assign_schedule(t)
    flash('Manual Championship pairing saved. The full Championship and consolation tree has been recalculated from these pairings.')
    return redirect(url_for('tournament_view', tournament_id=t.id) + '#knockout-pairing')


@app.route('/phase/<int:phase_id>/rules', methods=['POST'])
def update_phase_rules(phase_id):
    phase = Phase.query.get_or_404(phase_id)
    phase.best_of = int(request.form.get('best_of', phase.best_of))
    phase.points_to = int(request.form.get('points_to', phase.points_to))
    db.session.commit()
    flash('Phase rules updated.')
    return redirect(url_for('tournament_view', tournament_id=phase.tournament_id))


@app.route('/match/<int:match_id>/result', methods=['POST'])
def update_match_result(match_id):
    m = Match.query.get_or_404(match_id)
    if m.tournament.system == 'group_knockout' and not m.bracket.startswith('group_') and not group_stage_complete(m.tournament_id):
        flash('Finish all Group Stage matches first. Championship and Placement matches unlock automatically afterwards.')
        anchor = request.form.get('anchor') or ''
        return redirect(url_for('tournament_view', tournament_id=m.tournament_id) + (f'#{anchor}' if anchor else ''))
    if not m.player1_id or not m.player2_id:
        flash('This match is not ready yet. Its participants are still being determined by earlier results.')
        anchor = request.form.get('anchor') or ''
        return redirect(url_for('tournament_view', tournament_id=m.tournament_id) + (f'#{anchor}' if anchor else ''))

    # New score-grid format. Legacy text input remains supported for imported
    # backups and older browsers.
    sets = []
    max_sets = m.phase.best_of if m.phase else 5
    for idx in range(1, max_sets + 1):
        a_raw = request.form.get(f'set_{idx}_p1', '').strip()
        b_raw = request.form.get(f'set_{idx}_p2', '').strip()
        if a_raw == '' and b_raw == '':
            continue
        try:
            a, b = int(a_raw), int(b_raw)
        except Exception:
            flash(f'Set {idx}: enter both scores as whole numbers.')
            anchor = request.form.get('anchor') or ''
            return redirect(url_for('tournament_view', tournament_id=m.tournament_id) + (f'#{anchor}' if anchor else ''))
        sets.append([a, b])
    if not sets and request.form.get('score') is not None:
        sets, _, _ = parse_score(request.form.get('score', '').strip())

    p1_sets = sum(1 for a, b in sets if a > b)
    p2_sets = sum(1 for a, b in sets if b > a)
    winner_side = 1 if p1_sets > p2_sets else (2 if p2_sets > p1_sets else None)

    m.score_json = json.dumps(sets, ensure_ascii=False)
    m.winner_id = None
    m.loser_id = None
    result_meta = meta_for(m)
    if sets:
        result_meta['user_result_entered'] = True
    else:
        result_meta.pop('user_result_entered', None)
    set_meta(m, result_meta)
    if winner_side == 1:
        m.winner_id = m.player1_id
        m.loser_id = m.player2_id
    elif winner_side == 2:
        m.winner_id = m.player2_id
        m.loser_id = m.player1_id
    db.session.commit()

    was_group = m.bracket.startswith('group_')
    propagate_all(m.tournament_id)
    # When the final Group Stage result is entered, automatically produce a
    # sensible initial draw: strongest vs weakest while avoiding group rematches.
    if was_group and group_stage_complete(m.tournament_id):
        smart_byes = apply_smart_byes_after_groups(m.tournament)
        if not smart_byes:
            apply_recommended_entry_pairing(m.tournament)
    assign_schedule(m.tournament, renumber=False)
    flash('Result saved.')
    anchor = request.form.get('anchor') or ''
    return redirect(url_for('tournament_view', tournament_id=m.tournament_id) + (f'#{anchor}' if anchor else ''))


@app.route('/tournament/<int:tournament_id>/round_pairing', methods=['POST'])
def update_round_pairing(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    if settings_for(t).get('pairing_mode', 'first_round_only') != 'every_round':
        flash('Later-round reseeding is disabled in Settings. The bracket now follows its fixed tree.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')
    match_ids = [int(x) for x in request.form.getlist('match_id') if str(x).isdigit()]
    if not match_ids:
        flash('No round selected.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')
    matches = [Match.query.get(mid) for mid in match_ids]
    if any(m is None or m.tournament_id != t.id for m in matches):
        flash('Invalid round pairing request.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')
    metas = [meta_for(m) for m in matches]
    section_depth = {(meta.get('section') or m.round_name, int(meta.get('depth', 0))) for m, meta in zip(matches, metas)}
    if len(section_depth) != 1:
        flash('All edited matches must belong to the same round.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')

    descendants = downstream_match_ids(t, matches)
    protected = matches + [Match.query.get(mid) for mid in descendants if Match.query.get(mid)]
    started = next((m for m in protected if meta_for(m).get('user_result_entered') is True), None)
    if started:
        flash(f'This draw is locked because Match #{started.order_index} already has a saved result. Clear downstream results before reseeding this round.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')

    current_ids = []
    for m in matches:
        if not m.player1_id or not m.player2_id:
            flash('This round is not ready yet.')
            return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')
        current_ids.extend([m.player1_id, m.player2_id])
    selected = []
    new_pairs = []
    try:
        for m in matches:
            p1 = int(request.form.get(f'match_{m.id}_p1', '0'))
            p2 = int(request.form.get(f'match_{m.id}_p2', '0'))
            selected.extend([p1, p2])
            new_pairs.append((m, p1, p2))
    except Exception:
        flash('Invalid pairing selection.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')
    if sorted(selected) != sorted(current_ids) or len(set(selected)) != len(selected):
        flash('Use every eligible team exactly once in this round.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')

    # Avoid Group Stage rematches whenever a rematch-free pairing exists.
    groups = player_group_map(t)
    selected_rematches = sum(1 for _m, p1, p2 in new_pairs if groups.get(p1) and groups.get(p1) == groups.get(p2))
    recommended = recommended_pairs(t, current_ids)
    minimum_rematches = sum(1 for a, b in recommended if groups.get(a.id) and groups.get(a.id) == groups.get(b.id)) if recommended else selected_rematches
    if selected_rematches > minimum_rematches:
        flash('This draw contains an avoidable rematch between teams from the same Group Stage group.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')

    # Clear descendants so the changed draw propagates cleanly.
    for mid in descendants:
        dm = Match.query.get(mid)
        if not dm:
            continue
        dmeta = meta_for(dm)
        dm.player1_id = None
        dm.player2_id = None
        dm.score_json = '[]'
        dm.winner_id = None
        dm.loser_id = None
        dmeta.pop('user_result_entered', None)
        set_meta(dm, dmeta)

    for m, p1, p2 in new_pairs:
        meta = meta_for(m)
        meta['sources'] = [{'kind': 'player', 'player_id': p1}, {'kind': 'player', 'player_id': p2}]
        meta['manual_pairing'] = True
        set_meta(m, meta)
        m.player1_id = p1
        m.player2_id = p2
        m.score_json = '[]'
        m.winner_id = None
        m.loser_id = None
    db.session.commit()
    propagate_all(t.id)
    assign_schedule(t, renumber=False)
    flash('Round pairing updated. Downstream brackets were recalculated.')
    return redirect(url_for('tournament_view', tournament_id=t.id) + '#round-pairings')


@app.route('/tournament/<int:tournament_id>/logo', methods=['GET'])
def tournament_logo(tournament_id):
    t=Tournament.query.get_or_404(tournament_id); path=logo_path_for(t)
    if not path:
        return ('',404)
    return send_file(path)


@app.route('/tournament/<int:tournament_id>/logo', methods=['POST'])
def upload_tournament_logo(tournament_id):
    t=Tournament.query.get_or_404(tournament_id)
    f=request.files.get('logo')
    if not f or not f.filename:
        flash('Choose a PNG, JPG or WEBP logo file.')
        return redirect(url_for('tournament_view',tournament_id=t.id)+'#settings')
    ext=os.path.splitext(secure_filename(f.filename))[1].lower()
    if ext not in ('.png','.jpg','.jpeg','.webp'):
        flash('Logo must be PNG, JPG, JPEG or WEBP.')
        return redirect(url_for('tournament_view',tournament_id=t.id)+'#settings')
    filename=f'tournament_{t.id}_{uuid.uuid4().hex[:10]}{ext}'
    f.save(os.path.join(UPLOAD_DIR,filename))
    st=settings_for(t); old=st.get('logo_filename',''); st['logo_filename']=filename; save_settings_dict(t,st); db.session.commit()
    if old:
        try: os.remove(os.path.join(UPLOAD_DIR,os.path.basename(old)))
        except OSError: pass
    flash('Tournament logo uploaded.')
    return redirect(url_for('tournament_view',tournament_id=t.id)+'#settings')


@app.route('/tournament/<int:tournament_id>/logo/delete', methods=['POST'])
def delete_tournament_logo(tournament_id):
    t=Tournament.query.get_or_404(tournament_id); st=settings_for(t); old=st.get('logo_filename',''); st['logo_filename']=''; save_settings_dict(t,st); db.session.commit()
    if old:
        try: os.remove(os.path.join(UPLOAD_DIR,os.path.basename(old)))
        except OSError: pass
    flash('Tournament logo removed.')
    return redirect(url_for('tournament_view',tournament_id=t.id)+'#settings')


@app.route('/tournament/<int:tournament_id>/schedule', methods=['POST'])
def reschedule(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    reorder = request.form.get('reorder') == '1'
    if reorder and any(score_sets(m.score_json) or meta_for(m).get('user_result_entered') for m in t.matches if meta_for(m).get('key')):
        flash('Kolejność można zmienić przed wpisaniem wyników fazy pucharowej.')
        return redirect(url_for('tournament_view', tournament_id=t.id) + '#matches')
    t.court_count = max(1,int(request.form.get('court_count', t.court_count)))
    db.session.commit()
    assign_schedule(t, renumber=reorder)
    flash('Zapisano przeplataną kolejność meczów; mecz o 3. miejsce i finał zamykają harmonogram.' if reorder else 'Court allocation updated. Published match numbers were preserved.')
    return redirect(url_for('tournament_view', tournament_id=t.id))


@app.route('/tournament/<int:tournament_id>/export/db')
def export_db(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    buf = io.BytesIO(export_database_json(t).encode('utf-8'))
    return send_file(buf, as_attachment=True, download_name=f'{t.name}_backup.json', mimetype='application/json')


@app.route('/import/db', methods=['POST'])
def import_db():
    f = request.files.get('file')
    if not f:
        flash('No file selected.')
        return redirect(url_for('index'))
    t = import_database_json(f.read().decode('utf-8'))
    flash('Tournament database imported.')
    return redirect(url_for('tournament_view', tournament_id=t.id))


@app.route('/tournament/<int:tournament_id>/export/pdf')
def export_pdf(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    buf = write_bracket_pdf(t)
    return send_file(buf, as_attachment=True, download_name=f'{t.name}_bracket.pdf', mimetype='application/pdf')


@app.route('/tournament/<int:tournament_id>/export/groups_pdf')
def export_groups_pdf(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    buf = write_groups_pdf(t)
    return send_file(buf, as_attachment=True, download_name=f'{t.name}_groups.pdf', mimetype='application/pdf')


@app.route('/tournament/<int:tournament_id>/export/schedule_pdf')
def export_schedule_pdf(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    buf = write_schedule_pdf(t)
    return send_file(buf, as_attachment=True, download_name=f'{t.name}_schedule.pdf', mimetype='application/pdf')


@app.route('/tournament/<int:tournament_id>/export/final_standings_pdf')
def export_final_standings_pdf(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    buf = write_final_standings_pdf(t)
    return send_file(buf, as_attachment=True, download_name=f'{t.name}_final_standings.pdf', mimetype='application/pdf')


@app.route('/tournament/<int:tournament_id>/export/excel')
def export_excel(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    buf = write_results_excel(t)
    return send_file(buf, as_attachment=True, download_name=f'{t.name}_results.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/tournament/<int:tournament_id>/delete', methods=['POST'])
def delete_tournament(tournament_id):
    t = Tournament.query.get_or_404(tournament_id)
    db.session.delete(t)
    db.session.commit()
    flash('Tournament deleted.')
    return redirect(url_for('index'))


from online import configure_online
configure_online(app, db, Tournament, settings_for, save_settings_dict, tournament_view)
with app.app_context():
    repaired = repair_unplayed_brackets()
    upgrade_unplayed_schedules()
    if repaired:
        app.logger.info('Repaired pending bracket outcomes for tournament IDs: %s', repaired)

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(debug=False, use_reloader=False, port=APP_PORT)
